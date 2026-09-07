from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "project_progress_report_editable.pptx"
ASSETS = ROOT / "ppt_assets"
RESULT_PATH = ROOT / "runs/docking-loop-codex-gpt56-20260902-171109/real/result.json"
RESULT = json.loads(RESULT_PATH.read_text(encoding="utf-8"))["result"]
HISTORY = RESULT["docking_history"]
BEST_ATTEMPT = RESULT["best_attempt"]
BEST = next(x for x in HISTORY if x["attempt"] == BEST_ATTEMPT)

# ---------- Theme ----------
NAVY = RGBColor(18, 38, 63)
BLUE = RGBColor(39, 105, 171)
TEAL = RGBColor(27, 156, 142)
ORANGE = RGBColor(241, 139, 56)
RED = RGBColor(205, 77, 79)
GREEN = RGBColor(53, 153, 90)
PURPLE = RGBColor(116, 91, 166)
INK = RGBColor(34, 44, 56)
MID = RGBColor(91, 105, 120)
LIGHT = RGBColor(237, 242, 247)
LIGHT_BLUE = RGBColor(232, 241, 250)
LIGHT_TEAL = RGBColor(229, 246, 243)
LIGHT_ORANGE = RGBColor(253, 242, 231)
LIGHT_RED = RGBColor(251, 236, 236)
WHITE = RGBColor(255, 255, 255)
FONT = "Microsoft YaHei"
FONT_EN = "Aptos"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def set_run(run, size=18, bold=False, color=INK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=INK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.05,
             fill=None, line=None, radius=False):
    if fill is None and line is None and not radius:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill or WHITE
        shape.line.color.rgb = line or fill or WHITE
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    set_run(r, size=size, bold=bold, color=color)
    return shape


def add_rich_text(slide, x, y, w, h, runs, fill=None, line=None, radius=False,
                  valign=MSO_ANCHOR.TOP, margin=0.12, align=PP_ALIGN.LEFT):
    shape = add_text(slide, x, y, w, h, "", fill=fill, line=line, radius=radius,
                     valign=valign, margin=margin)
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    for item in runs:
        r = p.add_run()
        r.text = item[0]
        set_run(r, size=item[1], bold=item[2], color=item[3])
    return shape


def add_bullets(slide, x, y, w, h, items, size=17, color=INK, bullet_color=None,
                spacing=5, fill=None, line=None):
    shape = add_text(slide, x, y, w, h, "", fill=fill, line=line, radius=bool(fill), margin=0.18)
    tf = shape.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(spacing)
        p.line_spacing = 1.12
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p._p.get_or_add_pPr().insert(0, p._p._new_buChar()) if False else None
        # Use a visible bullet glyph to keep compatibility across PowerPoint versions.
        p.text = "• " + item
    return shape


def add_header(slide, title, section=None):
    add_text(slide, 0.55, 0.28, 11.9, 0.48, title, size=25, bold=True, color=NAVY,
             valign=MSO_ANCHOR.MIDDLE)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.84), Inches(1.0), Inches(0.06)).fill.solid()
    bar = slide.shapes[-1]
    bar.fill.fore_color.rgb = TEAL
    bar.line.color.rgb = TEAL
    if section:
        add_text(slide, 11.3, 0.33, 1.4, 0.3, section, size=10, bold=True, color=TEAL,
                 align=PP_ALIGN.RIGHT)


def add_footer(slide, page):
    add_text(slide, 0.55, 7.14, 10.5, 0.2,
             "Simple Molecular Agent｜项目进展汇报", size=8.5, color=MID)
    add_text(slide, 12.15, 7.12, 0.6, 0.22, str(page), size=9, color=MID,
             align=PP_ALIGN.RIGHT)


def add_connector(slide, x1, y1, x2, y2, color=MID, width=1.8, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        try:
            c.line.end_arrowhead = True
        except Exception:
            pass
    return c


def add_kpi(slide, x, y, w, h, value, label, color=TEAL, note=None):
    s = add_text(slide, x, y, w, h, "", fill=WHITE, line=LIGHT, radius=True, margin=0.16)
    tf = s.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = value; set_run(r, size=28, bold=True, color=color, font=FONT_EN)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r = p2.add_run(); r.text = label; set_run(r, size=12, bold=True, color=INK)
    if note:
        p3 = tf.add_paragraph(); p3.alignment = PP_ALIGN.CENTER
        r = p3.add_run(); r.text = note; set_run(r, size=8.5, color=MID)
    return s


def add_code_box(slide, x, y, w, h, text, title=None, font_size=11.5):
    if title:
        add_text(slide, x, y, w, 0.42, title, size=14, bold=True, color=NAVY)
        y += 0.46
        h -= 0.46
    shape = add_text(slide, x, y, w, h, "", fill=RGBColor(246, 248, 250),
                     line=RGBColor(218, 224, 230), radius=True, margin=0.18)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.line_spacing = 1.0
    r = p.add_run()
    r.text = text
    set_run(r, size=font_size, color=INK, font="Consolas")
    return shape


def style_table(table, header_color=NAVY, body_fill=WHITE, font_size=11):
    for c in range(len(table.columns)):
        cell = table.cell(0, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = header_color
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                set_run(r, size=font_size, bold=True, color=WHITE)
    for r_idx in range(1, len(table.rows)):
        for c in range(len(table.columns)):
            cell = table.cell(r_idx, c)
            cell.fill.solid(); cell.fill.fore_color.rgb = body_fill if r_idx % 2 else LIGHT
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    set_run(run, size=font_size, color=INK)


# ---------- Slide 1: Title ----------
slide = prs.slides.add_slide(blank)
bg = slide.background.fill
bg.solid(); bg.fore_color.rgb = NAVY
# decorative editable molecular-network circles
for x, y, d, col in [(10.2, 0.6, 1.2, TEAL), (11.5, 1.2, 0.55, ORANGE), (9.8, 2.0, 0.38, BLUE),
                      (11.0, 2.55, 0.8, PURPLE), (12.0, 3.1, 0.32, TEAL), (10.0, 4.0, 0.55, ORANGE)]:
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = col; shp.line.color.rgb = col
for a, b in [((10.8,1.2),(11.75,1.48)),((10.8,1.2),(10.0,2.15)),((10.0,2.15),(11.4,2.95)),
             ((11.4,2.95),(12.16,3.26)),((11.4,2.95),(10.28,4.28))]:
    add_connector(slide,*a,*b,color=RGBColor(105,145,173),width=2.0,arrow=False)
add_text(slide, 0.78, 0.78, 8.7, 0.42, "PROJECT PROGRESS REPORT", size=13, bold=True, color=TEAL)
add_text(slide, 0.78, 1.52, 8.9, 1.05, "LLM驱动的蛋白质–配体\n闭环分子优化智能体", size=34, bold=True, color=WHITE)
add_text(slide, 0.82, 3.02, 7.8, 0.75,
         "LLM提出假设｜Host确定性验证｜多随机种子Docking反馈｜全流程审计", size=17, color=RGBColor(205,220,232))
add_text(slide, 0.82, 5.52, 5.8, 0.36, "汇报人：__________", size=15, color=WHITE)
add_text(slide, 0.82, 6.02, 5.8, 0.36, "日期：__________", size=15, color=WHITE)
add_text(slide, 9.45, 6.55, 3.15, 0.32, "Simple Molecular Agent", size=12, bold=True, color=RGBColor(205,220,232), align=PP_ALIGN.RIGHT)

# ---------- Slide 2: Background ----------
slide = prs.slides.add_slide(blank); add_header(slide, "01 研究问题与项目目标", "BACKGROUND")
add_text(slide, 0.65, 1.15, 4.1, 0.48, "核心问题", size=19, bold=True, color=NAVY)
add_bullets(slide, 0.65, 1.68, 5.55, 3.8, [
    "分子优化高度依赖药物化学经验，搜索空间巨大",
    "LLM可提出建议，但可能产生错误原子、非法价态或空间碰撞",
    "单次Docking具有随机性，仅看单一分数容易误判",
    "设计过程缺少证据链时，难以复现“为什么生成该分子”",
], size=16, fill=LIGHT_BLUE, line=LIGHT_BLUE)
add_text(slide, 6.55, 1.15, 3.0, 0.48, "本项目的回答", size=19, bold=True, color=NAVY)
# central proposition
add_rich_text(slide, 6.55, 1.72, 5.95, 1.25, [
    ("LLM", 23, True, ORANGE), (" 负责提出化学假设\n", 18, True, INK),
    ("Host", 23, True, TEAL), (" 负责确定性事实与安全验证", 18, True, INK),
], fill=WHITE, line=TEAL, radius=True, valign=MSO_ANCHOR.MIDDLE, margin=0.24)
add_text(slide, 6.55, 3.25, 5.95, 0.58, "目标：建立可执行、可验证、可恢复、可审计的分子优化闭环", size=18, bold=True, color=WHITE, fill=NAVY, line=NAVY, radius=True, valign=MSO_ANCHOR.MIDDLE, margin=0.18)
add_text(slide, 6.55, 4.15, 1.75, 1.15, "输入\n共晶复合物PDB", size=15, bold=True, color=NAVY, fill=LIGHT, line=LIGHT, radius=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_text(slide, 8.73, 4.15, 1.75, 1.15, "过程\n闭环设计与筛选", size=15, bold=True, color=NAVY, fill=LIGHT_TEAL, line=LIGHT_TEAL, radius=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_text(slide, 10.9, 4.15, 1.6, 1.15, "输出\n候选与审计", size=15, bold=True, color=NAVY, fill=LIGHT_ORANGE, line=LIGHT_ORANGE, radius=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_connector(slide, 8.3, 4.72, 8.7, 4.72, color=TEAL, width=2.2)
add_connector(slide, 10.5, 4.72, 10.88, 4.72, color=TEAL, width=2.2)
add_text(slide, 0.65, 6.15, 11.85, 0.55, "当前边界：已实现到GNINA Docking；RBFE暂缓；结果属于计算排序信号，不等同于实验活性。", size=14, bold=True, color=RED, fill=LIGHT_RED, line=LIGHT_RED, radius=True, valign=MSO_ANCHOR.MIDDLE, margin=0.16)
add_footer(slide, 2)

# ---------- Slide 3: Architecture ----------
slide = prs.slides.add_slide(blank); add_header(slide, "02 系统总体架构", "ARCHITECTURE")
add_text(slide, 0.72, 1.10, 11.9, 0.42,
         "系统分成四层：先把结构读准确，再让模型提出方案，随后由程序执行，最后用统一协议比较。",
         size=15, color=MID)
architecture = [
    ("01", "可靠输入", "ComplexContext / FragmentLibrary（见04）", "把共晶PDB、任务和本地片段整理成可信的结构与化学输入"),
    ("02", "决策与调度", "LLM / Workflow（见03、04A）", "模型返回结构化JSON；工作流检查动作、证据和当前搜索状态"),
    ("03", "确定性执行", "ToolRegistry / Editing（见04、04B）", "程序查询位点并真正构建候选，返回可审计的JSON结果"),
    ("04", "比较与留痕", "Docking Adapter / AgentState（见04B）", "候选与参考按相同seed比较，再把结果写回下一轮上下文"),
]
for i, (num, title, comp, desc) in enumerate(architecture):
    x = 0.72 + i * 3.03
    add_text(slide, x, 1.78, 2.68, 0.52, f"{num}  {title}", size=17, bold=True, color=WHITE,
             fill=NAVY, line=NAVY, radius=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, x, 2.43, 2.68, 0.52, comp, size=12.5, bold=True, color=TEAL,
             fill=WHITE, line=LIGHT, radius=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, x, 3.08, 2.68, 1.72, desc, size=14, color=INK,
             fill=LIGHT, line=LIGHT, radius=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.16)
    if i < 3:
        add_connector(slide, x + 2.72, 3.60, x + 2.98, 3.60, color=TEAL, width=2.0)
add_text(slide, 1.0, 5.35, 11.3, 0.75,
         "分工边界：LLM只提出“查什么、改什么、为什么”；真正的结构事实、分子构建和评分由程序完成。",
         size=17, bold=True, color=NAVY, fill=LIGHT_TEAL, line=LIGHT_TEAL, radius=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.16)
add_footer(slide, 3)

# ---------- Slide 4: Loop ----------
slide = prs.slides.add_slide(blank); add_header(slide, "03 一次优化循环是怎样完成的", "WORKFLOW")
add_text(slide, 0.72, 1.08, 11.8, 0.42,
         "下面用非技术语言说明：系统每一轮都在回答六个连续问题。", size=15, color=MID)
workflow_steps = [
    ("1", "现在有什么？", "读取共晶结构\n确认原始配体和口袋"),
    ("2", "哪里可以改？", "查询局部环境、相互作用\n和可利用空间"),
    ("3", "准备怎么改？", "LLM选择位点、编辑方式\n和一个具体片段"),
    ("4", "这个分子做得出来吗？", "程序真正构建候选\n检查价态、电荷和碰撞"),
    ("5", "它是否比参考更好？", "候选和参考使用相同\nGNINA参数与三个seed"),
    ("6", "下一轮学到了什么？", "记录分数、pose和接触变化\n反馈给LLM继续判断"),
]
positions = [(0.72,1.72),(4.55,1.72),(8.38,1.72),(0.72,4.05),(4.55,4.05),(8.38,4.05)]
for (num, question, answer), (x, y) in zip(workflow_steps, positions):
    add_text(slide, x, y, 0.55, 0.55, num, size=17, bold=True, color=WHITE,
             fill=TEAL, line=TEAL, radius=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, x+0.68, y, 3.0, 0.48, question, size=16, bold=True, color=NAVY,
             valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, x, y+0.68, 3.55, 1.12, answer, size=13.5, color=INK,
             fill=WHITE, line=LIGHT, radius=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
# restrained flow connectors
add_connector(slide, 4.30, 2.82, 4.50, 2.82, color=MID, width=1.6)
add_connector(slide, 8.13, 2.82, 8.33, 2.82, color=MID, width=1.6)
add_connector(slide, 10.15, 3.60, 10.15, 3.98, color=MID, width=1.6)
add_connector(slide, 8.33, 5.15, 8.13, 5.15, color=MID, width=1.6)
add_connector(slide, 4.50, 5.15, 4.30, 5.15, color=MID, width=1.6)
add_text(slide, 0.85, 6.34, 11.65, 0.43,
         "循环的核心不是“让LLM直接画分子”，而是让每个方案都经过同一套可执行检查和对照评价。",
         size=14.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
add_footer(slide, 4)

# ---------- Slide 5: Implementation ----------
slide = prs.slides.add_slide(blank); add_header(slide, "04 核心模块：每一部分具体做什么", "IMPLEMENTATION")
add_text(slide, 0.72, 1.04, 11.85, 0.45,
         "代码按职责拆分。汇报时无需介绍函数细节，只需说明每个模块接收什么、解决什么问题、产生什么结果。",
         size=14.5, color=MID)
rows = [
    ("结构准备", "structure.py", "把完整PDB拆成蛋白和配体，并恢复正确的键级与芳香性", "可信的原始配体、受体和原子编号"),
    ("知识查询", "tools.py\nfragment_library.py", "回答“附近有什么残基、哪里有空间、有哪些可用片段”等具体问题", "Host事实与片段记录，不直接给优化结论"),
    ("决策与调度", "llm.py\nworkflow.py", "LLM提出查询和改造；Workflow检查证据、管理位点、去重并决定下一步", "一个有理由、可执行的具体transformation"),
    ("候选生成", "editing.py", "按指定原子或replacement site真正修改分子，并保留未修改骨架坐标", "候选SDF；失败时给出明确的化学或碰撞原因"),
    ("评价与记录", "adapters.py\nmodels.py", "候选与参考进行相同协议的多seed Docking，并保存状态和全部运行文件", "相对分数、pose一致性、接触变化和审计记录"),
]
# column headers
headers = [(0.72,1.65,1.65,"环节"),(2.45,1.65,2.05,"对应代码"),(4.58,1.65,5.15,"它解决什么问题"),(9.82,1.65,2.78,"输出什么")]
for x,y,w,text in headers:
    add_text(slide,x,y,w,0.5,text,size=14,bold=True,color=WHITE,fill=NAVY,line=NAVY,
             align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
for i, (stage, files, purpose, output) in enumerate(rows):
    y = 2.22 + i * 0.88
    fill = WHITE if i % 2 == 0 else LIGHT
    add_text(slide,0.72,y,1.65,0.78,stage,size=14,bold=True,color=NAVY,fill=fill,line=LIGHT,
             align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    add_text(slide,2.45,y,2.05,0.78,files,size=11.5,bold=True,color=TEAL,fill=fill,line=LIGHT,
             align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    add_text(slide,4.58,y,5.15,0.78,purpose,size=12.3,color=INK,fill=fill,line=LIGHT,
             valign=MSO_ANCHOR.MIDDLE,margin=0.14)
    add_text(slide,9.82,y,2.78,0.78,output,size=11.7,color=INK,fill=fill,line=LIGHT,
             valign=MSO_ANCHOR.MIDDLE,margin=0.12)
add_text(slide, 0.92, 6.68, 11.45, 0.34,
         "最重要的模块是 Workflow：它把模型、化学工具和Docking连接成一个有状态、可恢复的实验流程。",
         size=13.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
add_footer(slide, 5)

# ---------- Slide 6: JSON decisions ----------
slide = prs.slides.add_slide(blank); add_header(slide, "04A 模块之间如何通信：LLM返回的JSON", "JSON INTERFACE")
add_text(slide, 0.72, 1.02, 11.8, 0.42,
         "LLM不直接返回一个分子文件，而是返回固定格式的动作。Workflow读取 action 后决定调用工具还是进入候选验证。",
         size=14.5, color=MID)
query_json = '''{
  "action": "QUERY",
  "tool": "get_atom_environment",
  "arguments": {
    "atom_index": 9,
    "radius": 4.0
  },
  "question": "查看 atom 9 周围的蛋白环境"
}'''
ready_json = '''{
  "action": "READY",
  "parent_attempt": 15,
  "operation": "replace_hydrogen",
  "edit_atom_index": 9,
  "fragment_id": "chembl-brics-filtered-001052",
  "fragment_smiles": "C1CN(O[*:1])CCO1",
  "understanding": "已读取位点、空间和片段证据……",
  "edit_hypothesis": "用含氧杂环替换现有取代基……"
}'''
add_code_box(slide, 0.72, 1.60, 5.55, 3.92, query_json, "示例1：继续查询事实", font_size=12)
add_code_box(slide, 6.53, 1.60, 6.05, 3.92, ready_json, "示例2：提交一个具体改造", font_size=10.8)
add_text(slide, 0.82, 5.82, 5.35, 0.62,
         "QUERY 的含义：证据还不够，先向Host询问一个明确问题。",
         size=13.5, bold=True, color=NAVY, fill=LIGHT, line=LIGHT, radius=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_text(slide, 6.63, 5.82, 5.85, 0.62,
         "READY 的含义：方案已具体化，但仍需通过Host安全检查。",
         size=13.5, bold=True, color=NAVY, fill=LIGHT_TEAL, line=LIGHT_TEAL, radius=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_text(slide, 1.05, 6.62, 11.2, 0.28,
         "JSON让模型输出可以被程序校验、拒绝、重试和完整记录，而不是依赖自然语言猜测。",
         size=13, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
add_footer(slide, 6)

# ---------- Slide 7: JSON observations ----------
slide = prs.slides.add_slide(blank); add_header(slide, "04B Host返回什么：几何结果与Docking反馈", "JSON INTERFACE")
add_text(slide, 0.72, 1.02, 11.8, 0.42,
         "工具结果同样使用JSON。左侧回答“这个候选能否构建”，右侧回答“构建后在统一Docking协议下表现如何”。",
         size=14.5, color=MID)
geometry_json = '''{
  "status": "accepted",
  "canonical_smiles":
    "c1cc(... )cc(C2CC2)c1",
  "formal_charge": 0,
  "severe_clash_count": 0,
  "property_delta": {
    "heavy_atoms": 3,
    "molecular_weight": 40.06,
    "logp": 0.87,
    "tpsa": 0.0
  },
  "limitation": "This is not docking."
}'''
docking_json = '''{
  "attempt": 20,
  "design_region": "atom:9",
  "primary_metric": "minimizedAffinity",
  "delta_candidate_minus_reference": -1.7936,
  "seed_win_fraction": 1.0,
  "seed_stddev": 1.1421,
  "pose_consensus": {
    "mean_pairwise_rmsd": 0.373,
    "stable": true
  }
}'''
add_code_box(slide, 0.72, 1.58, 5.65, 4.58, geometry_json, "validate_candidate_geometry 返回示例", font_size=11.3)
add_code_box(slide, 6.63, 1.58, 5.95, 4.58, docking_json, "Docking反馈摘要（Attempt 20）", font_size=11.5)
add_text(slide, 0.86, 6.35, 11.55, 0.48,
         "数据流：LLM动作JSON → Host工具结果JSON → Candidate/Docking事件 → 压缩后进入下一轮LLM上下文。",
         size=14.5, bold=True, color=NAVY, fill=LIGHT, line=LIGHT, radius=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, 7)

# ---------- Slide 8: System boundary ----------
slide = prs.slides.add_slide(blank); add_header(slide, "05 系统边界：LLM负责什么，不负责什么", "BOUNDARY")
add_text(slide, 0.72, 1.03, 11.8, 0.42,
         "边界原则只有一句话：LLM负责提出可讨论的化学假设，Host负责决定该假设是否可以执行和如何评价。",
         size=15, color=MID)
# simple two-column layout without card grids
add_text(slide, 0.85, 1.72, 5.25, 0.48, "LLM负责", size=20, bold=True, color=NAVY)
add_bullets(slide, 0.85, 2.25, 5.25, 2.75, [
    "决定下一步需要查询什么事实",
    "基于现有证据选择位点、片段和编辑方式",
    "说明为什么提出该候选，以及还有哪些不确定性",
    "阅读Docking反馈后继续搜索、调整假设或请求停止",
], size=15.5, spacing=10)
# divider
line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.48), Inches(1.72), Inches(0.018), Inches(3.58))
line.fill.solid(); line.fill.fore_color.rgb = RGBColor(215,221,227); line.line.color.rgb = RGBColor(215,221,227)
add_text(slide, 6.88, 1.72, 5.25, 0.48, "LLM不负责", size=20, bold=True, color=NAVY)
add_bullets(slide, 6.88, 2.25, 5.25, 2.75, [
    "不能把自然语言建议直接当成有效分子",
    "不能自由修改原子编号、切割键或删除原子集合",
    "不能跳过价态、电荷、碰撞、去重和证据检查",
    "不能把Docking分数解释成实验活性或真实自由能",
], size=15.5, spacing=10)
add_text(slide, 0.85, 5.32, 11.3, 0.54,
         "Host的硬边界：结构事实、合法编辑、候选构建、是否进入Docking以及最终审计，全部由确定性程序控制。",
         size=15.5, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
add_text(slide, 1.12, 6.08, 10.8, 0.58,
         "LLM建议  ≠  候选已接受  ≠  Docking表现更好  ≠  实验活性提高",
         size=19, bold=True, color=NAVY, fill=LIGHT, line=LIGHT, radius=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, 8)

# ---------- Slide 9: Progress KPIs ----------
slide = prs.slides.add_slide(blank); add_header(slide, "06 当前工程进展", "PROGRESS")
add_kpi(slide,0.65,1.28,2.25,1.45,"16","受控工具",NAVY,"结构、位点、片段、几何")
add_kpi(slide,3.1,1.28,2.25,1.45,"20,039","统一片段",TEAL,"ChEMBL + curated，去重后")
add_kpi(slide,5.55,1.28,2.25,1.45,"100 / 100","自动化测试通过",NAVY,"molecular-agent-docking环境")
add_kpi(slide,8.0,1.28,2.25,1.45,"12,953","Python代码行",TEAL,"核心模块 + 脚本 + 测试")
add_kpi(slide,10.45,1.28,2.25,1.45,"3 seeds","配对Docking",NAVY,"17 / 29 / 43")
add_text(slide, 0.65, 3.05, 4.1, 0.48, "已完成能力", size=19, bold=True, color=NAVY)
add_bullets(slide,0.65,3.55,5.75,2.6,[
    "端到端：PDB → 候选 → GNINA → 反馈",
    "多seed参考基线与候选配对比较",
    "候选、pose、命令和日志完整落盘",
    "支持中断恢复且不跨运行注入候选",
],size=14.5,fill=LIGHT_TEAL,line=LIGHT_TEAL,spacing=7)
add_text(slide, 6.75, 3.05, 4.1, 0.48, "当前正在增强", size=19, bold=True, color=NAVY)
add_bullets(slide,6.75,3.55,5.95,2.6,[
    "结构化工作记忆与LLM上下文压缩",
    "parent-specific局部优化与多代候选",
    "请求payload/token用量记录",
    "Elite archive与候选集合优化策略",
],size=14.5,fill=LIGHT_ORANGE,line=LIGHT_ORANGE,spacing=7)
add_footer(slide, 9)

# ---------- Slide 10: Latest run ----------
slide = prs.slides.add_slide(blank); add_header(slide, "07 代表性真实运行：按替换位点看局部趋势", "RESULTS")
add_text(slide, 0.72, 1.02, 8.3, 0.42,
         "本次运行实际搜索了三个不同原子位点。不同位点不能串成一条趋势线，因此分别展示各位点内部的搜索过程。",
         size=14, color=MID)
add_text(slide, 9.35, 0.98, 3.25, 0.48,
         "78次Docking = atom 9: 30｜atom 6: 38｜atom 3: 10",
         size=11.5, bold=True, color=NAVY, fill=LIGHT, line=LIGHT, radius=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
# common legend
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.35), Inches(1.51), Inches(0.42), Inches(0.035)).fill.solid()
slide.shapes[-1].fill.fore_color.rgb = RGBColor(145,158,171); slide.shapes[-1].line.color.rgb = RGBColor(145,158,171)
add_text(slide,4.82,1.39,1.7,0.25,"当次候选",size=10.5,color=MID)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.15), Inches(1.51), Inches(0.42), Inches(0.05)).fill.solid()
slide.shapes[-1].fill.fore_color.rgb = TEAL; slide.shapes[-1].line.color.rgb = TEAL
add_text(slide,6.62,1.39,2.4,0.25,"该位点 best-so-far",size=10.5,color=MID)

site_specs = [
    ("atom:9", "Atom 9", 0.55),
    ("atom:6", "Atom 6", 4.47),
    ("atom:3", "Atom 3", 8.39),
]
for site_key, site_title, x in site_specs:
    site_history = [item for item in HISTORY if item.get("design_region") == site_key]
    local_best_quality = None
    local_best_delta = None
    local_best_series = []
    local_best_item = None
    for item in site_history:
        quality = item.get("quality")
        if item.get("stability_eligible") and quality is not None and (local_best_quality is None or quality > local_best_quality):
            local_best_quality = quality
            local_best_delta = item.get("delta_candidate_minus_reference")
            local_best_item = item
        local_best_series.append(local_best_delta)
    cd = CategoryChartData()
    cd.categories = [str(i+1) for i in range(len(site_history))]
    cd.add_series("当次候选", [item.get("delta_candidate_minus_reference") for item in site_history])
    cd.add_series("位点best-so-far", local_best_series)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x), Inches(1.78), Inches(3.72), Inches(3.62), cd).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = f"{site_title} 局部搜索"
    chart.has_legend = False
    chart.value_axis.minimum_scale = -2.1
    chart.value_axis.maximum_scale = 1.2
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = LIGHT
    chart.category_axis.tick_labels.font.size = Pt(8)
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.series[0].format.line.color.rgb = RGBColor(145,158,171)
    chart.series[0].format.line.width = Pt(1.3)
    chart.series[1].format.line.color.rgb = TEAL
    chart.series[1].format.line.width = Pt(2.6)
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.name = FONT
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(13)
    eligible_count = sum(bool(item.get("stability_eligible")) for item in site_history)
    pose_rmsd = (local_best_item.get("pose_consensus") or {}).get("mean_pairwise_rmsd") if local_best_item else None
    best_attempt = local_best_item.get("attempt") if local_best_item else None
    best_delta = local_best_item.get("delta_candidate_minus_reference") if local_best_item else None
    summary = (
        f"Docking {len(site_history)} 次｜通过seed门槛 {eligible_count} 次\n"
        f"位点最佳：Attempt {best_attempt}｜Δ = {best_delta:.3f}｜pose RMSD = {pose_rmsd:.3f} Å"
        if local_best_item and pose_rmsd is not None else
        f"Docking {len(site_history)} 次｜通过seed门槛 {eligible_count} 次"
    )
    add_text(slide, x, 5.52, 3.72, 0.78, summary, size=10.8, bold=True, color=NAVY,
             fill=LIGHT, line=LIGHT, radius=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0.10)
add_text(slide,0.75,6.48,11.85,0.35,
         "结论：Atom 9给出本次全局最佳；Atom 6有中等改善；Atom 3虽有评分改善，但最佳pose RMSD较高，稳定性需要谨慎解释。",
         size=13.2,bold=True,color=NAVY,align=PP_ALIGN.CENTER)
add_footer(slide, 10)

# ---------- Slide 11: Best candidate ----------
slide = prs.slides.add_slide(blank); add_header(slide, "08 最佳候选：结构、性质与多seed结果", "BEST CANDIDATE")
# molecule images
slide.shapes.add_picture(str(ASSETS/"reference_ligand.png"), Inches(0.55), Inches(1.05), width=Inches(3.2), height=Inches(1.72))
slide.shapes.add_picture(str(ASSETS/"best_candidate.png"), Inches(3.78), Inches(1.05), width=Inches(3.2), height=Inches(1.72))
add_text(slide,0.65,2.62,6.25,0.32,"Attempt 20｜atom 9 局部优化｜generation 3｜新增含氧杂环片段",size=12.5,bold=True,color=NAVY,align=PP_ALIGN.CENTER)
# property table
rows,cols=7,3
shape=slide.shapes.add_table(rows,cols,Inches(7.18),Inches(1.10),Inches(5.58),Inches(2.25))
t=shape.table
data=[("性质","参考配体","最佳候选"),("分子量","323.40","424.50"),("LogP","4.06","3.68"),("HBD / HBA","2 / 5","2 / 8"),("TPSA","75.72","97.42"),("重原子","24","31"),("可旋转键","5","7")]
for r,row in enumerate(data):
    for c,val in enumerate(row): t.cell(r,c).text=val
style_table(t,font_size=10.5)
# per-seed editable column chart
metrics=BEST["comparison"]["per_seed"]
cd=CategoryChartData(); cd.categories=[str(x["seed"]) for x in metrics]
cd.add_series("候选",[float(x["candidate_pose"]["properties"]["minimizedAffinity"]) for x in metrics])
cd.add_series("参考",[float(x["reference_pose"]["properties"]["minimizedAffinity"]) for x in metrics])
chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(0.65),Inches(3.32),Inches(6.25),Inches(3.25),cd).chart
chart.has_title=True; chart.chart_title.text_frame.text="三个seed的 minimizedAffinity（越低越好）"
chart.has_legend=True; chart.legend.position=XL_LEGEND_POSITION.BOTTOM
chart.series[0].format.fill.solid(); chart.series[0].format.fill.fore_color.rgb=TEAL
chart.series[1].format.fill.solid(); chart.series[1].format.fill.fore_color.rgb=RGBColor(160,175,190)
chart.value_axis.has_major_gridlines=True; chart.value_axis.major_gridlines.format.line.color.rgb=LIGHT
chart.chart_title.text_frame.paragraphs[0].runs[0].font.name=FONT; chart.chart_title.text_frame.paragraphs[0].runs[0].font.size=Pt(14)
# summary metrics
add_text(slide,7.18,3.62,2.58,1.12,"-1.794\n平均相对差",size=22,bold=True,color=WHITE,fill=TEAL,line=TEAL,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
add_text(slide,10.02,3.62,2.74,1.12,"3 / 3\nseed方向一致",size=22,bold=True,color=WHITE,fill=NAVY,line=NAVY,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
add_text(slide,7.18,4.98,2.58,1.12,"0.373 Å\n平均pose RMSD",size=20,bold=True,color=NAVY,fill=LIGHT,line=LIGHT,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
add_text(slide,10.02,4.98,2.74,1.12,"+0.769\nCNNaffinity均值差",size=20,bold=True,color=NAVY,fill=LIGHT_TEAL,line=LIGHT_TEAL,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
add_text(slide,7.18,6.28,5.58,0.38,"新增共识接触：ASN132 / GLU12 / GLY11；丢失：GLU81 / LYS89",size=11.5,bold=True,color=RED,align=PP_ALIGN.CENTER)
add_footer(slide, 11)

# ---------- Slide 12: Interpretation ----------
slide = prs.slides.add_slide(blank); add_header(slide, "09 如何解释当前结果", "INTERPRETATION")
add_text(slide,0.65,1.25,5.8,0.58,"当前证据支持的结论",size=19,bold=True,color=WHITE,fill=GREEN,line=GREEN,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
add_bullets(slide,0.65,1.95,5.8,3.9,[
    "端到端闭环已实际跑通，并具有自动化测试和审计证据",
    "最佳候选在固定GNINA协议下主指标平均优于参考",
    "三个seed的主指标方向一致，rank-1 pose具有较好一致性",
    "系统能够保留历史最佳，而非只输出最后一次尝试",
],size=15,fill=LIGHT_TEAL,line=LIGHT_TEAL,spacing=9)
add_text(slide,6.85,1.25,5.8,0.58,"当前不能直接得出的结论",size=19,bold=True,color=WHITE,fill=RED,line=RED,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
add_bullets(slide,6.85,1.95,5.8,3.9,[
    "不能称为实验结合活性提高1.79 kcal/mol",
    "不能证明真实结合自由能已经改善",
    "不能证明搜索已科学收敛：运行因硬安全上限停止",
    "不能证明方法已泛化：目前主要验证1H1Q单体系",
    "不能直接将Docking最佳候选等同于RBFE或合成候选",
],size=15,fill=LIGHT_RED,line=LIGHT_RED,spacing=8)
add_text(slide,0.9,6.12,11.55,0.54,"推荐表述：获得了“固定Docking协议下具有进一步验证价值的计算候选”，而不是“已经证明活性提高”。",size=16,bold=True,color=NAVY,fill=LIGHT_ORANGE,line=LIGHT_ORANGE,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, 12)

# ---------- Slide 13: Next steps ----------
slide = prs.slides.add_slide(blank); add_header(slide, "10 下一阶段计划", "NEXT STEPS")
road=[
    ("01","结果复核","重复独立Docking\n共晶pose重现\n关键相互作用复查",NAVY),
    ("02","候选集合","Elite archive\n结构聚类与多样性\n性质/pose联合排序",TEAL),
    ("03","方法对照","工具预算消融\n单seed vs 多seed\nparent优化对照",NAVY),
    ("04","高精度验证","原子映射与参数化\nRBFE cohort/network\n实验验证",TEAL),
]
for i,(num,title,desc,c) in enumerate(road):
    x=0.72+i*3.12
    circle=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x+0.82),Inches(1.38),Inches(0.72),Inches(0.72))
    circle.fill.solid(); circle.fill.fore_color.rgb=c; circle.line.color.rgb=c
    add_text(slide,x+0.82,1.38,0.72,0.72,num,size=16,bold=True,color=WHITE,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    add_text(slide,x,2.25,2.38,0.55,title,size=18,bold=True,color=WHITE,fill=c,line=c,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    add_text(slide,x,2.93,2.38,2.0,desc,size=14,color=INK,fill=WHITE,line=LIGHT,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    if i<3: add_connector(slide,x+2.47,3.45,x+3.02,3.45,color=MID,width=2.0)
add_text(slide,0.72,5.42,11.9,0.9,"目标转变：从“寻找单一高分分子”升级为“在固定计算预算内构建高质量、稳定且具有结构多样性的候选集合”。",size=19,bold=True,color=WHITE,fill=NAVY,line=NAVY,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE,margin=0.16)
add_text(slide,1.15,6.55,11.0,0.34,"Docking elite archive → RBFE shortlist → RBFE cohorts / perturbation networks → 实验",size=14,bold=True,color=TEAL,align=PP_ALIGN.CENTER)
add_footer(slide, 13)

# ---------- Slide 14: Evidence paths ----------
slide = prs.slides.add_slide(blank); add_header(slide, "附录：可复核材料与文件位置", "APPENDIX")
items=[
    ("流程说明","workflow_explanation.md / workflow.png",NAVY),
    ("最新完整结果","runs/docking-loop-codex-gpt56-20260902-171109/real/result.json",TEAL),
    ("Docking历史",".../real/docking-history.json",NAVY),
    ("最佳候选",".../real/candidate-20.sdf",TEAL),
    ("参考配体",".../real/reference-ligand.sdf",NAVY),
    ("项目说明","README.md / LLM_BATCHED_MULTISITE_OPTIMIZATION.md",TEAL),
]
for i,(name,path,c) in enumerate(items):
    y=1.25+i*0.83
    add_text(slide,0.75,y,2.15,0.58,name,size=14,bold=True,color=WHITE,fill=c,line=c,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    add_text(slide,3.05,y,9.55,0.58,path,size=13,color=INK,fill=WHITE,line=LIGHT,radius=True,valign=MSO_ANCHOR.MIDDLE,margin=0.16)
add_text(slide,0.85,6.42,11.55,0.5,"本PPT中的文本、流程框、表格和图表均可在PowerPoint/WPS中直接编辑；分子结构作为高分辨率图片插入。",size=14,bold=True,color=NAVY,fill=LIGHT_BLUE,line=LIGHT_BLUE,radius=True,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, 14)

# Metadata
prs.core_properties.title = "LLM驱动的蛋白质-配体闭环分子优化智能体：项目进展汇报"
prs.core_properties.subject = "Simple Molecular Agent project progress"
prs.core_properties.author = "OpenAI coding assistant"
prs.core_properties.keywords = "LLM, molecular design, GNINA, docking, RDKit, editable PowerPoint"
prs.core_properties.comments = "Editable slides generated from current project code and run results."

prs.save(OUT)
print(OUT)
