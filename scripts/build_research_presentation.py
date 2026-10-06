from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE

ROOT = Path(__file__).resolve().parents[1]
ASSET = ROOT / "assets" / "ppt"
OUTDIR = ROOT / "deliverables"
OUTDIR.mkdir(parents=True, exist_ok=True)
OUT = OUTDIR / "科研汇报_保留骨架的分子优化智能体.pptx"

# 16:9 canvas
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# Palette
NAVY = RGBColor(7, 20, 38)
NAVY2 = RGBColor(13, 35, 59)
TEAL = RGBColor(20, 184, 166)
CYAN = RGBColor(56, 189, 248)
AMBER = RGBColor(242, 163, 58)
INK = RGBColor(27, 45, 65)
MUTED = RGBColor(82, 98, 115)
LIGHT = RGBColor(245, 248, 252)
BORDER = RGBColor(220, 229, 239)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(16, 148, 126)
RED = RGBColor(210, 83, 83)
FONT = "WenQuanYi Zen Hei"
FONT_LATIN = "Aptos"


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, x, y, w, h, text, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT,
             margin=0.0, linesp=1.0, rotation=0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.rotation = rotation
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = linesp
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_text(slide, x, y, w, h, parts, size=18, color=INK,
                  valign=MSO_ANCHOR.TOP, margin=0.0, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for part in parts:
        run = p.add_run()
        run.text = part.get("text", "")
        run.font.name = part.get("font", FONT)
        run.font.size = Pt(part.get("size", size))
        run.font.bold = part.get("bold", False)
        run.font.color.rgb = part.get("color", color)
    return box


def add_rect(slide, x, y, w, h, fill, radius=True, line=None, transparency=0):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    return shp


def add_circle(slide, x, y, d, fill, line=None, transparency=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.fill.transparency = transparency
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
    return shp


def add_line(slide, x1, y1, x2, y2, color=MUTED, width=1.5, dash=None, arrow_end=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if dash:
        line.line.dash_style = dash
    if arrow_end:
        line.line.end_arrowhead = True
    return line


def add_picture_crop(slide, path, x, y, w, h):
    path = str(path)
    im = Image.open(path)
    iw, ih = im.size
    target = w / h
    source = iw / ih
    pic = slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    if source > target:
        shown = target / source
        crop = (1 - shown) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    elif source < target:
        shown = source / target
        crop = (1 - shown) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop
    return pic


def add_picture_contain(slide, path, x, y, w, h):
    path = str(path)
    im = Image.open(path)
    iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    return slide.shapes.add_picture(path, Inches(px), Inches(py), width=Inches(pw), height=Inches(ph))


def add_title(slide, number, title, subtitle=None, dark=False):
    col = WHITE if dark else INK
    muted = RGBColor(177, 201, 223) if dark else MUTED
    add_text(slide, 0.60, 0.34, 0.55, 0.32, f"{number:02d}", 11, AMBER, True, valign=MSO_ANCHOR.MIDDLE)
    add_line(slide, 1.16, 0.50, 1.55, 0.50, AMBER, 2.2)
    add_text(slide, 1.70, 0.24, 10.9, 0.52, title, 25, col, True, valign=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, 1.70, 0.76, 10.7, 0.34, subtitle, 10.5, muted, False, valign=MSO_ANCHOR.MIDDLE)


def add_footer(slide, text, dark=False, page=None):
    color = RGBColor(146, 168, 190) if dark else RGBColor(120, 137, 154)
    add_text(slide, 0.62, 7.14, 11.7, 0.20, text, 7.2, color, False, valign=MSO_ANCHOR.MIDDLE)
    if page is not None:
        add_text(slide, 12.30, 7.10, 0.38, 0.22, str(page), 8, color, True, align=PP_ALIGN.RIGHT)


def add_tag(slide, x, y, text, fill=TEAL, color=WHITE, width=None):
    width = width or max(0.9, 0.19 * len(text) + 0.32)
    add_rect(slide, x, y, width, 0.34, fill, radius=True)
    add_text(slide, x, y+0.01, width, 0.30, text, 9, color, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_metric_card(slide, x, y, w, number, label, accent=TEAL, note=None):
    add_rect(slide, x, y, w, 1.00, WHITE, radius=True, line=BORDER)
    add_rect(slide, x, y, 0.09, 1.00, accent, radius=False)
    add_text(slide, x+0.25, y+0.12, w-0.35, 0.40, number, 22, accent, True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, x+0.25, y+0.53, w-0.35, 0.25, label, 10, INK, True)
    if note:
        add_text(slide, x+0.25, y+0.77, w-0.35, 0.17, note, 7.3, MUTED)


# ---------------------------------------------------------------------------
# Slide 1 — cover
slide = prs.slides.add_slide(BLANK)
add_picture_crop(slide, ASSET / "molecular_agent_cover.png", 0, 0, 13.333, 7.5)
add_rect(slide, 0, 0, 6.35, 7.5, NAVY, radius=False)
add_tag(slide, 0.72, 0.62, "科研汇报 · 外行友好版", fill=TEAL, width=2.35)
add_text(slide, 0.72, 1.36, 5.45, 1.42, "让 AI 帮忙改分子，\n但不让它自由发挥", 31, WHITE, True, linesp=0.93)
add_text(slide, 0.74, 3.05, 5.10, 0.92,
         "一个保留已知药物骨架、结合蛋白口袋信息、\n并由计算工具逐步验真的分子优化智能体", 15, RGBColor(211, 228, 241), False, linesp=1.12)
add_line(slide, 0.74, 4.18, 2.05, 4.18, AMBER, 3.0)
add_text(slide, 0.74, 4.42, 4.9, 0.72,
         "核心思路：LLM 提假设，Host 做化学与结构验证，\nDocking 提供反馈，再进入下一轮。", 12, WHITE, False)
add_tag(slide, 0.74, 6.48, "共晶结构", fill=NAVY2, width=1.20)
add_tag(slide, 2.06, 6.48, "单点编辑", fill=NAVY2, width=1.20)
add_tag(slide, 3.38, 6.48, "多 seed 验证", fill=NAVY2, width=1.55)
add_text(slide, 10.50, 6.73, 2.18, 0.25, "Simple Molecular Agent", 8.5, RGBColor(203, 221, 237), True, align=PP_ALIGN.RIGHT)

# ---------------------------------------------------------------------------
# Slide 2 — problem
slide = prs.slides.add_slide(BLANK)
set_bg(slide, LIGHT)
add_title(slide, 2, "它要解决什么问题？", "药物分子优化不是“分数越高越好”，而是多重约束下的可解释搜索")

# Left intro statement
add_rect(slide, 0.66, 1.30, 6.12, 0.78, NAVY, radius=True)
add_rich_text(slide, 0.96, 1.45, 5.55, 0.45, [
    {"text": "同一骨架只改一小段，", "color": WHITE, "bold": True, "size": 16},
    {"text": "结果却可能完全不同", "color": AMBER, "bold": True, "size": 16},
], valign=MSO_ANCHOR.MIDDLE)

challenges = [
    ("01", "选择空间很大", "一个位点就可能有几十到成千上万种侧链；\n真正昂贵的是后续构建与 docking。"),
    ("02", "能接上 ≠ 能结合", "化学价态正确还不够：要放得进、方向对、\n还要保留关键骨架与锚点。"),
    ("03", "单次评分会骗人", "不同随机 seed、不同 pose 可能给出不同答案；\n必须看稳定性与一致性。"),
]
for i, (num, head, body) in enumerate(challenges):
    y = 2.32 + i*1.27
    add_rect(slide, 0.66, y, 6.12, 1.03, WHITE, radius=True, line=BORDER)
    add_circle(slide, 0.90, y+0.20, 0.54, TEAL if i < 2 else AMBER)
    add_text(slide, 0.90, y+0.20, 0.54, 0.54, num, 10, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 1.62, y+0.14, 2.20, 0.28, head, 13.5, INK, True)
    add_text(slide, 1.62, y+0.48, 4.78, 0.43, body, 9.3, MUTED)

# Right visual
add_rect(slide, 7.05, 1.30, 5.62, 5.39, WHITE, radius=True, line=BORDER)
add_picture_crop(slide, ASSET / "molecular_agent_cover.png", 7.17, 1.42, 5.38, 3.82)
add_rect(slide, 7.45, 4.76, 4.82, 1.52, NAVY, radius=True, transparency=3)
add_text(slide, 7.76, 4.98, 4.20, 0.34, "通俗类比：钥匙配锁", 15, AMBER, True)
add_text(slide, 7.76, 5.38, 4.20, 0.65,
         "不只是“能塞进去”，还要方向正确、\n关键齿位不变，而且多次尝试都能复现。", 10.5, WHITE)
add_footer(slide, "问题边界：当前目标是计算筛选与排序，不直接等同于实验活性或临床效果。", page=2)

# ---------------------------------------------------------------------------
# Slide 3 — what the project does
slide = prs.slides.add_slide(BLANK)
set_bg(slide, WHITE)
add_title(slide, 3, "当前项目具体做什么？", "以共晶结构为锚点，对已知配体做“保留骨架的单点微调”")

# molecule panels
add_rect(slide, 0.70, 1.38, 5.05, 3.36, LIGHT, radius=True, line=BORDER)
add_rect(slide, 7.58, 1.38, 5.05, 3.36, LIGHT, radius=True, line=BORDER)
add_tag(slide, 0.96, 1.62, "起点", fill=NAVY2, width=0.83)
add_text(slide, 1.93, 1.59, 2.8, 0.32, "原始共晶配体", 13, INK, True)
add_tag(slide, 7.84, 1.62, "候选", fill=TEAL, width=0.83)
add_text(slide, 8.80, 1.59, 3.2, 0.32, "单一位点的类似物", 13, INK, True)
add_picture_contain(slide, ASSET / "reference_ligand.png", 0.98, 2.06, 4.48, 2.28)
add_picture_contain(slide, ASSET / "best_candidate.png", 7.86, 2.06, 4.48, 2.28)

# central transformation
add_circle(slide, 6.05, 2.27, 1.23, NAVY)
add_text(slide, 6.05, 2.27, 1.23, 0.55, "单点", 14, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.BOTTOM)
add_text(slide, 6.05, 2.83, 1.23, 0.43, "侧链替换", 9.5, AMBER, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.TOP)
add_line(slide, 5.56, 2.88, 6.00, 2.88, TEAL, 2.3, arrow_end=True)
add_line(slide, 7.31, 2.88, 7.54, 2.88, TEAL, 2.3, arrow_end=True)

add_rect(slide, 0.70, 4.95, 11.93, 1.44, NAVY, radius=True)
concepts = [
    ("01", "骨架保留", "不做无约束 scaffold hopping"),
    ("02", "口袋条件化", "利用局部空间与残基环境"),
    ("03", "相互作用感知", "关注锚点、gain / loss"),
    ("04", "完整可审计", "每个决策、结构、分数可追溯"),
]
for i, (n, h, b) in enumerate(concepts):
    x = 0.98 + i*2.92
    add_circle(slide, x, 5.24, 0.48, TEAL if i < 3 else AMBER)
    add_text(slide, x, 5.24, 0.48, 0.48, n, 8.5, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, x+0.63, 5.17, 1.90, 0.28, h, 12, WHITE, True)
    add_text(slide, x+0.63, 5.54, 2.00, 0.42, b, 8.1, RGBColor(184, 207, 226))
add_rect(slide, 3.19, 6.56, 6.94, 0.42, RGBColor(232, 248, 245), radius=True)
add_text(slide, 3.19, 6.58, 6.94, 0.36,
         "不是让 LLM 随机写 SMILES，而是让它在受控动作空间里做策略选择。", 10, GREEN, True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, "示例结构来自 4WKQ–gefitinib 闭集 benchmark 的本地运行产物。", page=3)

# ---------------------------------------------------------------------------
# Slide 4 — architecture
slide = prs.slides.add_slide(BLANK)
set_bg(slide, NAVY)
add_title(slide, 4, "项目框架：LLM 做策略，Host 做裁判", "核心是一个“提出假设 → 确定性验证 → 反馈学习”的闭环", dark=True)

# Input block
add_rect(slide, 0.62, 1.46, 2.15, 4.74, NAVY2, radius=True, line=RGBColor(42, 82, 111))
add_tag(slide, 0.88, 1.73, "输入", fill=AMBER, width=0.78)
input_items = [
    ("共晶复合物", "蛋白 + 原始配体三维坐标"),
    ("任务约束", "允许位点、预算、保护核心"),
    ("本地片段库", "真实 fragment ID 与性质"),
]
for i, (h, b) in enumerate(input_items):
    yy = 2.35 + i*1.08
    add_circle(slide, 0.90, yy+0.03, 0.25, TEAL)
    add_text(slide, 1.28, yy-0.01, 1.15, 0.26, h, 11, WHITE, True)
    add_text(slide, 0.90, yy+0.32, 1.54, 0.38, b, 7.8, RGBColor(174, 199, 221))

# Center system container
add_rect(slide, 3.08, 1.46, 7.23, 4.74, RGBColor(10, 28, 49), radius=True, line=RGBColor(42, 82, 111))
add_tag(slide, 3.36, 1.72, "智能体闭环", fill=TEAL, width=1.20)

# Dossier
add_rect(slide, 3.42, 2.22, 1.55, 2.92, RGBColor(22, 48, 72), radius=True, line=RGBColor(54, 101, 132))
add_text(slide, 3.63, 2.46, 1.10, 0.36, "结构 dossier", 12, WHITE, True, align=PP_ALIGN.CENTER)
for i, t in enumerate(["ligand 图", "pocket 环境", "interaction", "合法位点", "片段面板"]):
    add_rect(slide, 3.68, 3.00+i*0.38, 1.00, 0.27, RGBColor(31, 67, 93), radius=True)
    add_text(slide, 3.68, 3.00+i*0.38, 1.00, 0.27, t, 7.6, RGBColor(209, 230, 243), align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)

# LLM
add_rect(slide, 5.38, 2.22, 2.02, 1.16, RGBColor(13, 87, 92), radius=True, line=TEAL)
add_text(slide, 5.62, 2.44, 1.54, 0.30, "LLM 策略层", 13, WHITE, True, align=PP_ALIGN.CENTER)
add_text(slide, 5.58, 2.83, 1.64, 0.28, "QUERY / READY / STOP", 7.4, RGBColor(184, 244, 235), align=PP_ALIGN.CENTER)

# Host stages
host_stages = [
    (5.28, 4.03, 1.40, "RDKit 构建", "价态 · 电荷 · 去重"),
    (6.90, 4.03, 1.40, "几何预筛", "碰撞 · 空间 · 身份"),
    (8.52, 4.03, 1.48, "Docking / PLIP", "pose · 锚点 · 多 seed"),
]
for x, y, w, h, b in host_stages:
    add_rect(slide, x, y, w, 1.12, RGBColor(26, 52, 77), radius=True, line=RGBColor(57, 104, 137))
    add_text(slide, x+0.08, y+0.17, w-0.16, 0.28, h, 10.3, WHITE, True, align=PP_ALIGN.CENTER)
    add_text(slide, x+0.08, y+0.58, w-0.16, 0.28, b, 7.3, RGBColor(176, 203, 225), align=PP_ALIGN.CENTER)

# arrows within loop
add_line(slide, 4.98, 2.80, 5.33, 2.80, TEAL, 2.1, arrow_end=True)
add_line(slide, 6.38, 3.40, 6.00, 3.98, CYAN, 2.0, arrow_end=True)
add_line(slide, 6.69, 4.59, 6.86, 4.59, CYAN, 2.0, arrow_end=True)
add_line(slide, 8.31, 4.59, 8.49, 4.59, CYAN, 2.0, arrow_end=True)
# feedback path
add_line(slide, 9.27, 3.99, 9.27, 3.55, AMBER, 2.0)
add_line(slide, 9.27, 3.55, 7.64, 3.55, AMBER, 2.0)
add_line(slide, 7.64, 3.55, 7.64, 2.80, AMBER, 2.0, arrow_end=True)
add_text(slide, 7.93, 3.25, 1.10, 0.22, "压缩反馈", 8, AMBER, True, align=PP_ALIGN.CENTER)
add_text(slide, 5.58, 5.45, 3.80, 0.30, "Host 决定“能不能做、是否可信”", 10, RGBColor(184, 207, 226), True, align=PP_ALIGN.CENTER)

# Output block
add_rect(slide, 10.62, 1.46, 2.10, 4.74, NAVY2, radius=True, line=RGBColor(42, 82, 111))
add_tag(slide, 10.90, 1.73, "输出", fill=AMBER, width=0.78)
output_items = [
    ("历史最佳候选", "不是最后一次，而是最佳且稳定者"),
    ("位点级 SAR", "知道什么改动有效 / 无效"),
    ("完整审计", "JSON、SDF、命令、pose、日志"),
]
for i, (h, b) in enumerate(output_items):
    yy = 2.38 + i*1.08
    add_text(slide, 10.92, yy, 1.45, 0.26, h, 10.5, WHITE, True)
    add_text(slide, 10.92, yy+0.34, 1.45, 0.40, b, 7.5, RGBColor(174, 199, 221))

add_line(slide, 2.79, 3.80, 3.03, 3.80, AMBER, 2.4, arrow_end=True)
add_line(slide, 10.33, 3.80, 10.58, 3.80, AMBER, 2.4, arrow_end=True)
add_rect(slide, 3.92, 6.46, 5.45, 0.42, RGBColor(16, 54, 72), radius=True)
add_text(slide, 3.92, 6.48, 5.45, 0.36, "一次请求规划多个候选；宿主批量筛选，只把有价值的结果反馈给模型。", 9.4, RGBColor(207, 233, 241), True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, "当前主流程运行到 docking；RBFE 保留为未来阶段，不伪造自由能结果。", dark=True, page=4)

# ---------------------------------------------------------------------------
# Slide 5 — credibility / gates
slide = prs.slides.add_slide(BLANK)
set_bg(slide, LIGHT)
add_title(slide, 5, "为什么这个框架比“让 AI 直接给答案”更可信？", "关键不是模型有多会说，而是每个建议都要穿过确定性的证据门")

# Two-column responsibilities
add_rect(slide, 0.68, 1.35, 5.90, 2.27, WHITE, radius=True, line=BORDER)
add_rect(slide, 6.75, 1.35, 5.90, 2.27, WHITE, radius=True, line=BORDER)
add_tag(slide, 0.96, 1.62, "LLM：策略层", fill=TEAL, width=1.48)
add_tag(slide, 7.03, 1.62, "Host：事实与裁判", fill=NAVY2, width=1.75)
llm_items = ["选择值得尝试的位点与片段", "提出可检验的结构假设", "阅读趋势，决定继续、查询或停止"]
host_items = ["执行精确分子图编辑与化学校验", "做碰撞、pose、锚点和多 seed 门控", "保存结构、分数、命令和完整审计"]
for i, t in enumerate(llm_items):
    add_circle(slide, 1.00, 2.15+i*0.42, 0.18, TEAL)
    add_text(slide, 1.31, 2.08+i*0.42, 4.82, 0.30, t, 10.3, INK, i==0)
for i, t in enumerate(host_items):
    add_circle(slide, 7.08, 2.15+i*0.42, 0.18, NAVY2)
    add_text(slide, 7.39, 2.08+i*0.42, 4.82, 0.30, t, 10.3, INK, i==0)

# Evidence gates
add_text(slide, 0.72, 3.93, 4.3, 0.36, "四道证据门", 16, INK, True)
gates = [
    ("1", "化学合法", "价态 / 电荷 / 结构身份", TEAL),
    ("2", "几何可行", "空间碰撞 / 构象 / 去重", CYAN),
    ("3", "姿态保留", "固定核心 RMSD / 关键锚点", AMBER),
    ("4", "结果稳定", "多 seed 一致 / interaction", GREEN),
]
for i, (n, h, b, c) in enumerate(gates):
    x = 0.72 + i*3.06
    add_rect(slide, x, 4.40, 2.72, 1.55, WHITE, radius=True, line=BORDER)
    add_circle(slide, x+0.18, 4.65, 0.52, c)
    add_text(slide, x+0.18, 4.65, 0.52, 0.52, n, 12, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, x+0.85, 4.59, 1.62, 0.32, h, 12, INK, True)
    add_text(slide, x+0.24, 5.18, 2.23, 0.42, b, 8.6, MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        add_line(slide, x+2.74, 5.17, x+3.00, 5.17, c, 2.0, arrow_end=True)

add_rect(slide, 1.83, 6.28, 9.67, 0.57, NAVY, radius=True)
add_rich_text(slide, 1.83, 6.30, 9.67, 0.51, [
    {"text": "智能 ≠ 自由发挥；  ", "size": 12, "bold": True, "color": AMBER},
    {"text": "智能 = 提出假设 + 接受可复现的计算反馈", "size": 12, "bold": True, "color": WHITE},
], valign=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
add_footer(slide, "严格门控也会拒绝“看起来分数不错、但姿态不可信”的候选。", page=5)

# ---------------------------------------------------------------------------
# Slide 6 — current validation
slide = prs.slides.add_slide(BLANK)
set_bg(slide, WHITE)
add_title(slide, 6, "当前进展：链路已跑通，而且知道哪些结论还不能下", "4WKQ–EGFR / gefitinib 闭集侧链 benchmark 的本地审计结果")

add_metric_card(slide, 0.64, 1.25, 2.90, "3 / 3", "参考配体 seed 校准通过", TEAL, "核心 RMSD 0.297–0.413 Å")
add_metric_card(slide, 3.66, 1.25, 2.90, "78 / 78", "匿名侧链可精确构建", CYAN, "固定切口、图身份全部通过")
add_metric_card(slide, 6.68, 1.25, 2.90, "20 / 20", "LLM 候选完成 pose 评价", AMBER, "均通过本次 pose-retention 门")
add_metric_card(slide, 9.70, 1.25, 2.90, "16 / 20", "达到 seed 稳定性资格", GREEN, "best 在 3/3 seeds 改善")

# charts
add_rect(slide, 0.64, 2.48, 4.03, 3.72, LIGHT, radius=True, line=BORDER)
add_text(slide, 0.93, 2.68, 3.45, 0.32, "参考校准：姿态远低于 2 Å 门槛", 12, INK, True)
add_picture_contain(slide, ASSET / "reference_calibration.png", 0.84, 3.04, 3.62, 2.77)

add_rect(slide, 4.86, 2.48, 7.75, 3.72, LIGHT, radius=True, line=BORDER)
add_text(slide, 5.14, 2.68, 7.12, 0.32, "一次 LLM 运行中，20 个候选相对参考配体的 docking 差值", 12, INK, True)
add_picture_contain(slide, ASSET / "llm_candidate_scores.png", 5.08, 3.00, 7.28, 2.93)

# honest boundary
add_rect(slide, 0.64, 6.39, 11.97, 0.53, RGBColor(255, 246, 229), radius=True, line=RGBColor(244, 201, 130))
add_rich_text(slide, 0.88, 6.42, 11.50, 0.46, [
    {"text": "必须诚实：", "size": 9.5, "bold": True, "color": RGBColor(154, 91, 11)},
    {"text": "单次 LLM run 的 best Δ=-0.367；单次随机基线 best Δ=-0.507。", "size": 9.5, "bold": True, "color": INK},
    {"text": "  说明系统已可公平比较，但尚不能声称 LLM 优于基线，更不能声称实验活性提升。", "size": 9.5, "color": MUTED},
], valign=MSO_ANCHOR.MIDDLE)
add_footer(slide, "来源：4WKQ/BENCHMARK_RUNNING.md、参考 calibration.json、LLM 与 random baseline 的 result.json。", page=6)

# ---------------------------------------------------------------------------
# Slide 7 — expected impact and next steps
slide = prs.slides.add_slide(BLANK)
add_picture_crop(slide, ASSET / "molecular_agent_outcome.png", 0, 0, 13.333, 7.5)
add_rect(slide, 0, 0, 7.14, 7.5, NAVY, radius=False)
add_tag(slide, 0.70, 0.58, "预期效果", fill=AMBER, width=1.12)
add_text(slide, 0.70, 1.15, 5.95, 0.70, "把“盲目试错”变成\n可解释、可复现的候选收敛", 25, WHITE, True, linesp=0.95)

impacts = [
    ("更少浪费", "先批量规划和几何筛选，再把计算预算给更有价值的候选。"),
    ("更易解释", "每个候选都能追溯到位点、片段、假设、pose 与相互作用变化。"),
    ("更可复现", "固定参考、固定 seeds、完整日志，避免只展示一次“好看的分数”。"),
]
for i, (h, b) in enumerate(impacts):
    yy = 2.38 + i*0.96
    add_circle(slide, 0.76, yy, 0.46, TEAL if i < 2 else AMBER)
    add_text(slide, 0.76, yy, 0.46, 0.46, str(i+1), 10, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 1.41, yy-0.02, 1.46, 0.27, h, 12.5, WHITE, True)
    add_text(slide, 2.85, yy-0.03, 3.43, 0.47, b, 9.2, RGBColor(199, 219, 235))

add_text(slide, 0.72, 5.49, 2.0, 0.30, "下一步验证", 13, AMBER, True)
for i, t in enumerate(["多随机种子、多个模型与 score-guided baseline 的重复实验", "更可靠的局部 docking / RBFE 与不确定性分析", "对优先候选进行合成与湿实验验证"]):
    add_circle(slide, 0.79, 5.92+i*0.34, 0.12, RGBColor(170, 213, 233))
    add_text(slide, 1.06, 5.84+i*0.34, 5.43, 0.28, t, 8.8, WHITE)

add_rect(slide, 7.68, 5.67, 4.78, 1.03, RGBColor(5, 18, 34), radius=True, transparency=12)
add_text(slide, 7.98, 5.86, 4.18, 0.28, "最终交付物", 10.5, AMBER, True)
add_text(slide, 7.98, 6.18, 4.18, 0.35, "稳定、可解释的候选清单，而不是“自动证明的新药”。", 11.3, WHITE, True)
add_footer(slide, "一句话总结：让 LLM 负责“想”，让确定性化学与结构工具负责“验”。", dark=True, page=7)

# Metadata
prs.core_properties.title = "保留骨架、口袋条件化的分子优化智能体"
prs.core_properties.subject = "科研汇报：项目问题、框架、架构、当前进展与预期效果"
prs.core_properties.author = "Simple Molecular Agent 项目"
prs.core_properties.comments = "由项目工作区资料与 image-gen 生成的视觉资产自动构建。"

prs.save(OUT)
print(OUT)
