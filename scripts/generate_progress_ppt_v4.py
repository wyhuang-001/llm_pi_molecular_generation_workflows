from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "项目进展汇报_可编辑_v4.pptx"
RESULT_FILE = ROOT / "runs/docking-loop-codex-gpt56-20260902-171109/real/result.json"
ASSETS = ROOT / "ppt_assets"
RESULT = json.loads(RESULT_FILE.read_text(encoding="utf-8"))["result"]
HISTORY = RESULT["docking_history"]
BEST = next(x for x in HISTORY if x["attempt"] == RESULT["best_attempt"])

# Restrained research-presentation palette.
BLACK = RGBColor(30, 34, 38)
DARK = RGBColor(44, 55, 67)
BLUE = RGBColor(42, 91, 132)
TEAL = RGBColor(30, 126, 119)
GRAY = RGBColor(100, 108, 116)
LINE = RGBColor(210, 215, 220)
PALE = RGBColor(246, 247, 248)
PALE_BLUE = RGBColor(239, 244, 248)
WHITE = RGBColor(255, 255, 255)
RED = RGBColor(170, 65, 65)
GREEN = RGBColor(43, 121, 76)
PALE_RED = RGBColor(251, 239, 239)
PALE_GREEN = RGBColor(237, 247, 241)
FONT = "Microsoft YaHei"
CODE = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def set_run(run, size=16, bold=False, color=BLACK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def text(slide, x, y, w, h, value, size=16, bold=False, color=BLACK,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, fill=None, border=None,
         margin=0.05, font=FONT):
    if fill is None and border is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = fill or WHITE
        shape.line.color.rgb = border or fill or WHITE
        shape.line.width = Pt(0.8)
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value
    set_run(r, size, bold, color, font)
    return shape


def title(slide, value, page):
    text(slide, 0.62, 0.30, 11.9, 0.48, value, 24, True, BLACK, valign=MSO_ANCHOR.MIDDLE)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(0.88), Inches(12.05), Inches(0.025))
    rule.fill.solid(); rule.fill.fore_color.rgb = BLUE; rule.line.color.rgb = BLUE
    text(slide, 12.2, 7.12, 0.45, 0.20, str(page), 9, False, GRAY, align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, h, items, size=16, color=BLACK, gap=7):
    shape = text(slide, x, y, w, h, "", margin=0)
    tf = shape.text_frame; tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.line_spacing = 1.10
        r = p.add_run(); r.text = "• " + item
        set_run(r, size, False, color)
    return shape


def code_block(slide, x, y, w, h, value, heading=None, size=10.5):
    if heading:
        text(slide, x, y, w, 0.34, heading, 14, True, DARK)
        y += 0.39; h -= 0.39
    return text(slide, x, y, w, h, value, size, False, BLACK,
                fill=PALE, border=LINE, margin=0.15, font=CODE)


def connector(slide, x1, y1, x2, y2, color=GRAY, width=1.3):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    try: c.line.end_arrowhead = True
    except Exception: pass
    return c


def table_style(tbl, header=True, size=11):
    for r in range(len(tbl.rows)):
        for c in range(len(tbl.columns)):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.fill.solid(); cell.fill.fore_color.rgb = PALE_BLUE if r == 0 and header else (WHITE if r % 2 else PALE)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for run in p.runs:
                    set_run(run, size, r == 0 and header, DARK)


# 1. Title
s = prs.slides.add_slide(blank)
text(s, 0.78, 1.15, 11.7, 1.0, "蛋白质–配体分子优化工作流", 34, True, BLACK)
text(s, 0.80, 2.18, 11.0, 0.56, "当前实现、数据接口与初步Docking结果", 22, False, BLUE)
rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.80), Inches(3.05), Inches(2.2), Inches(0.035))
rule.fill.solid(); rule.fill.fore_color.rgb = BLUE; rule.line.color.rgb = BLUE
text(s, 0.80, 3.42, 8.8, 0.82,
     "项目核心：LLM提出下一步操作，程序查询结构、构建候选并完成统一协议的Docking比较。",
     18, False, DARK)
text(s, 0.80, 5.75, 4.5, 0.35, "汇报人：__________", 15, False, GRAY)
text(s, 0.80, 6.20, 4.5, 0.35, "日期：__________", 15, False, GRAY)
text(s, 10.1, 6.62, 2.4, 0.30, "Simple Molecular Agent", 11, False, GRAY, align=PP_ALIGN.RIGHT)

# 2. Problem
s = prs.slides.add_slide(blank); title(s, "01 研究问题与当前目标", 2)
text(s, 0.75, 1.20, 5.5, 0.40, "研究问题", 19, True, DARK)
bullets(s, 0.78, 1.72, 5.5, 3.55, [
    "如何从一个共晶配体出发，自动提出可执行的局部结构修改？",
    "如何避免LLM给出错误原子、非法价态、空间碰撞或重复候选？",
    "如何把Docking结果返回给模型，形成连续的局部优化过程？",
    "如何保存每次查询、候选和评分，便于复核和恢复？",
], 16)
text(s, 6.70, 1.20, 5.4, 0.40, "当前目标", 19, True, DARK)
bullets(s, 6.73, 1.72, 5.6, 3.55, [
    "输入：完整蛋白质–配体共晶结构",
    "设计：氢取代或Host定义的片段替换",
    "筛选：化学检查、三维碰撞检查和多seed GNINA",
    "输出：候选结构、相对评分、pose和完整运行记录",
], 16)
text(s, 0.78, 5.70, 11.55, 0.64,
     "当前阶段做到Docking。RBFE和实验验证尚未完成，因此结果只表示固定计算协议下的候选排序。",
     16, True, RED, fill=PALE, border=LINE, margin=0.15, valign=MSO_ANCHOR.MIDDLE)

# 3. Architecture with concrete payloads
s = prs.slides.add_slide(blank); title(s, "02 系统总体架构：查询反馈和Docking反馈都进入下一轮", 3)
modules = [
    (0.45, "1  读入任务和结构", "structure.py / fragment_library.py",
'''{
 "complex_path": "complex.pdb",
 "ligand_selector": {"chain":"A", "residue_name":"2A6"},
 "fragment_library_path": "fragments_unified.json"
}'''),
    (3.62, "2  模型发出QUERY或READY", "llm.py / workflow.py",
'''{
 "action": "QUERY",
 "tool": "get_atom_environment",
 "arguments": {"atom_index":9, "radius":4.0}
}'''),
    (6.79, "3  程序执行工具并返回结果", "tools.py / editing.py",
'''{
 "atom_index": 9,
 "replaceable_hydrogens": 1,
 "protein_atoms": [
   {"atom":"HIS:A:84:O", "distance":3.048}, ...
 ]
}'''),
    (9.96, "4  构建候选并完成Docking", "editing.py / adapters.py / models.py",
'''{
 "geometry_status": "accepted",
 "delta_candidate_minus_reference": -1.7936,
 "seed_win_fraction": 1.0,
 "pose_rmsd": 0.373
}'''),
]
for x, name, files, payload in modules:
    text(s, x, 1.28, 2.90, 0.48, name, 15.5, True, BLACK)
    text(s, x, 1.78, 2.90, 0.30, files, 10.0, False, BLUE, font=CODE)
    code_block(s, x, 2.17, 2.90, 2.70, payload, size=8.8)
# Forward path
connector(s,3.39,3.48,3.57,3.48,BLUE,1.5)
connector(s,6.56,3.48,6.74,3.48,BLUE,1.5)
connector(s,9.73,3.48,9.91,3.48,BLUE,1.5)
text(s,5.82,2.90,0.70,0.24,"QUERY",9.5,True,BLUE,align=PP_ALIGN.CENTER)
text(s,8.84,2.90,0.90,0.24,"READY后",9.5,True,BLUE,align=PP_ALIGN.CENTER)
# Query feedback loop: module 3 -> module 2
connector(s,8.25,5.02,8.25,5.42,GRAY,1.2)
connector(s,8.25,5.42,5.02,5.42,GRAY,1.2)
connector(s,5.02,5.42,5.02,5.02,BLUE,1.5)
text(s,5.34,5.12,2.70,0.24,"工具结果返回模型，可继续下一次查询",10.5,True,BLUE,align=PP_ALIGN.CENTER)
# Docking feedback loop: module 4 -> module 2
connector(s,11.42,5.02,11.42,6.13,GRAY,1.2)
connector(s,11.42,6.13,4.65,6.13,GRAY,1.2)
connector(s,4.65,6.13,4.65,5.02,BLUE,1.5)
text(s,7.25,5.82,3.40,0.24,"Docking结果返回模型，开始下一轮设计",10.5,True,BLUE,align=PP_ALIGN.CENTER)
text(s,0.75,6.58,11.8,0.28,
     "因此有两个循环：QUERY ↔ 工具反馈可以重复多轮；READY通过后进入构建和Docking，再把评分反馈给模型。",
     13.5, True, DARK, align=PP_ALIGN.CENTER)

# 4. Workflow
s = prs.slides.add_slide(blank); title(s, "03 一轮分子优化的实际流程", 4)
steps = [
    ("1", "读取原始结构", "确定配体、蛋白和原子编号"),
    ("2", "查询一个位点", "获得附近残基、可替换氢和空间信息"),
    ("3", "查询片段", "从本地片段库筛选操作兼容的片段"),
    ("4", "提交具体修改", "指定操作、原子/切割位点和fragment"),
    ("5", "构建并检查候选", "RDKit构建；检查价态、电荷、碰撞和重复"),
    ("6", "与参考配体对接", "三个相同seed分别比较候选和参考"),
    ("7", "返回结果继续搜索", "反馈分数、pose和接触变化"),
]
for i,(n,t,d) in enumerate(steps):
    y=1.18+i*0.75
    text(s,0.78,y,0.42,0.42,n,15,True,WHITE,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE,fill=BLUE,border=BLUE)
    text(s,1.42,y-0.01,2.42,0.38,t,16,True,DARK)
    text(s,4.02,y-0.01,7.95,0.42,d,15,False,BLACK)
    if i<6:
        line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(0.98),Inches(y+0.44),Inches(0.012),Inches(0.28))
        line.fill.solid();line.fill.fore_color.rgb=LINE;line.line.color.rgb=LINE
text(s,0.78,6.55,11.5,0.34,"LLM负责步骤2、3、4和7中的判断；其余步骤由程序执行。",14,True,BLUE)

# 5. Code modules
s = prs.slides.add_slide(blank); title(s, "04 代码模块与实际职责", 5)
rows = [
    ("structure.py", "读取PDB并恢复配体化学结构", "配体RDKit分子、蛋白原子、receptor PDB"),
    ("fragment_library.py", "读取和检索本地片段库", "片段SMILES、性质、允许操作和来源"),
    ("tools.py", "执行模型请求的结构和片段查询", "ToolObservation JSON"),
    ("llm.py", "把当前状态发送给模型，并解析模型返回的JSON", "QUERY、READY、STOP等动作"),
    ("workflow.py", "管理查询顺序、证据、位点、候选历史和恢复", "下一步动作与最终result.json"),
    ("editing.py", "按指定transformation真正构建候选", "candidate SDF或明确失败原因"),
    ("adapters.py", "调用GNINA并解析三个seed结果", "每个seed的pose、score和相对统计"),
    ("models.py", "保存运行状态和工作记忆", "decisions、observations、history、memory"),
]
shape=s.shapes.add_table(len(rows)+1,3,Inches(0.72),Inches(1.28),Inches(11.9),Inches(5.55))
tbl=shape.table
for c,v in enumerate(["文件","主要职责","产生的结果"]): tbl.cell(0,c).text=v
tbl.columns[0].width=Inches(2.35);tbl.columns[1].width=Inches(5.0);tbl.columns[2].width=Inches(4.55)
for r,row in enumerate(rows,1):
    for c,v in enumerate(row): tbl.cell(r,c).text=v
table_style(tbl,True,11.5)

# 6. 04A query + feedback
s = prs.slides.add_slide(blank); title(s, "04A 查询工具示例：查看Atom 9周围环境", 6)
query='''{
  "action": "QUERY",
  "tool": "get_atom_environment",
  "arguments": {
    "atom_index": 9,
    "radius": 4.0
  }
}'''
feedback='''{
  "atom_index": 9,
  "parent_attempt": null,
  "element": "C",
  "aromatic": true,
  "replaceable_hydrogens": 1,
  "protein_atoms": [
    {"atom":"HIS:A:84:O",   "distance":3.048},
    {"atom":"ILE:A:10:CG2", "distance":3.404},
    {"atom":"HIS:A:84:C",   "distance":3.907},
    {"atom":"GLN:A:85:CA",  "distance":3.919}
  ]
}'''
code_block(s,0.75,1.28,4.85,4.55,query,"模型发出的工具调用",12)
connector(s,5.78,3.45,6.36,3.45,BLUE,1.8)
code_block(s,6.55,1.28,5.95,4.55,feedback,"工具返回的实际结果",10.7)
text(s,0.78,6.18,11.65,0.50,
     "这次返回说明：Atom 9是芳香碳，有1个可替换氢；4 Å内最近的是HIS84主链O和ILE10侧链。",
     15,True,DARK,fill=PALE_BLUE,border=LINE,margin=0.14,valign=MSO_ANCHOR.MIDDLE)

# 7. 04B batch + feedback
s = prs.slides.add_slide(blank); title(s, "04B 查询工具示例：批量预筛Atom 9的烷基片段", 7)
batch_query='''{
  "action": "QUERY",
  "tool": "generate_site_candidate_batch",
  "arguments": {
    "target_type": "atom",
    "target_id": 9,
    "size_class": "small",
    "chemical_tag": "alkyl",
    "limit": 8
  }
}'''
batch_feedback='''{
  "status": "complete",
  "target_id": 9,
  "library_match_count": 8,
  "accepted_count": 3,
  "rejected_count": 5,
  "accepted": [
    {"fragment_id":"...000446", "clashes":0},
    {"fragment_id":"...000547", "clashes":0},
    {"fragment_id":"...000550", "clashes":0}
  ],
  "rejected_example": {
    "fragment_id": "curated-ethyl",
    "failure_class": "steric_clash",
    "severe_clash_count": 1
  }
}'''
code_block(s,0.75,1.25,5.1,4.8,batch_query,"输入：限定同一位点、尺寸和化学类型",10.9)
connector(s,6.00,3.55,6.42,3.55,BLUE,1.8)
code_block(s,6.57,1.25,5.93,4.8,batch_feedback,"输出：8个库片段中3个通过几何预筛",10.4)
text(s,0.78,6.32,11.65,0.40,
     "批量工具只做片段检索、候选构建和碰撞预筛，不执行Docking。LLM再从通过的候选中选择一个继续。",
     14.5,True,DARK,align=PP_ALIGN.CENTER)

# 8. 04C ready -> geometry -> docking
s = prs.slides.add_slide(blank); title(s, "04C 从READY到Docking反馈", 8)
ready='''{
 "action": "READY",
 "parent_attempt": 15,
 "operation": "replace_hydrogen",
 "edit_atom_index": 9,
 "fragment_id": "...001052",
 "fragment_smiles": "C1CN(O[*:1])CCO1"
}'''
geometry='''{
 "status": "accepted",
 "formal_charge": 0,
 "severe_clash_count": 0,
 "molecular_weight": 424.50,
 "logp": 3.68,
 "hba": 8,
 "tpsa": 97.42
}'''
docking='''{
 "attempt": 20,
 "design_region": "atom:9",
 "delta_candidate_minus_reference": -1.7936,
 "seed_win_fraction": 1.0,
 "seed_stddev": 1.1421,
 "pose_consensus": {
   "mean_pairwise_rmsd": 0.373,
   "stable": true
 }
}'''
code_block(s,0.55,1.30,3.65,4.65,ready,"1. LLM提交具体修改",9.7)
connector(s,4.28,3.62,4.55,3.62,BLUE,1.5)
code_block(s,4.63,1.30,3.65,4.65,geometry,"2. 程序构建并检查",10.1)
connector(s,8.36,3.62,8.63,3.62,BLUE,1.5)
code_block(s,8.71,1.30,4.05,4.65,docking,"3. GNINA结果返回工作流",9.8)
text(s,0.78,6.25,11.65,0.46,
     "只有第2步accepted后才执行第3步；第3步结果会成为下一轮模型输入，而不是直接宣布候选有效。",
     14.5,True,DARK,fill=PALE,border=LINE,margin=0.12,valign=MSO_ANCHOR.MIDDLE)

# 9. Boundary
s = prs.slides.add_slide(blank); title(s, "05 系统边界", 9)
text(s,0.78,1.28,5.45,0.40,"LLM可以决定",19,True,DARK)
bullets(s,0.80,1.82,5.45,2.72,[
    "下一步查询哪个工具",
    "选择哪个Host支持的位点和片段",
    "提出具体的结构修改及理由",
    "根据Docking结果继续搜索或请求停止",
],16)
text(s,6.72,1.28,5.45,0.40,"LLM不能决定",19,True,DARK)
bullets(s,6.74,1.82,5.45,2.72,[
    "不能自由指定不存在的原子或切割键",
    "不能跳过化学、几何和重复检查",
    "不能直接修改或伪造Docking分数",
    "不能把Docking结果表述成实验活性",
],16)
line=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(6.43),Inches(1.28),Inches(0.015),Inches(3.45))
line.fill.solid();line.fill.fore_color.rgb=LINE;line.line.color.rgb=LINE
text(s,0.82,4.95,11.45,0.46,"程序必须保证：输入结构一致、编辑可执行、候选可复现、评分协议一致、失败原因有记录。",16,True,TEAL)
text(s,1.02,5.86,11.0,0.62,"LLM建议  ≠  候选通过检查  ≠  Docking更好  ≠  实验活性提高",19,True,BLACK,align=PP_ALIGN.CENTER,fill=PALE,border=LINE,margin=0.14,valign=MSO_ANCHOR.MIDDLE)

# 10. Implemented components
s = prs.slides.add_slide(blank); title(s, "06 当前已经实现的内容", 10)
implemented = [
    (0.78,1.25,"结构输入","从完整共晶PDB识别蛋白和配体；恢复配体键级、芳香性、电荷和氢数；生成reference ligand与protein-only receptor。"),
    (0.78,2.78,"LLM与工具调用","LLM返回QUERY、READY、STOP等JSON；程序提供16个结构、位点、片段和候选几何工具，并把结果作为JSON返回。"),
    (0.78,4.31,"候选分子构建","支持replace_hydrogen和replace_fragment；使用RDKit生成候选，保留未修改骨架坐标，并检查价态、电荷、碰撞和重复结构。"),
    (6.78,1.25,"本地片段库","统一片段库20,039条；每条记录包含SMILES、性质、尺寸、化学标签、允许操作和ChEMBL来源。"),
    (6.78,2.78,"多seed Docking","使用GNINA对候选和参考配体分别对接；固定seed为17、29、43；输出四类评分、seed稳定性和pose一致性。"),
    (6.78,4.31,"运行记录与恢复","保存每次LLM决策、工具反馈、候选SDF、Docking命令、日志和结果JSON；支持从同一运行目录恢复；当前100项测试通过。"),
]
for x,y,h,d in implemented:
    text(s,x,y,5.55,0.36,h,18,True,DARK)
    text(s,x,y+0.48,5.55,0.82,d,14.2,False,BLACK)
    rule=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y+1.28),Inches(5.45),Inches(0.012))
    rule.fill.solid();rule.fill.fore_color.rgb=LINE;rule.line.color.rgb=LINE
text(s,0.82,6.43,11.45,0.34,"以上均为当前代码和已有运行中可以直接复核的功能。",13.5,True,BLUE,align=PP_ALIGN.CENTER)

# 11. Fragment database source
s = prs.slides.add_slide(blank); title(s, "07 片段数据库来源与处理", 11)
text(s,0.78,1.18,3.1,0.40,"主要来源：ChEMBL 37",19,True,DARK)
bullets(s,0.80,1.68,5.65,2.15,[
    "官方FTP：chembl_37_chemreps.txt.gz",
    "下载并处理895,532个分子；原始文件共2,897,819行",
    "许可：CC BY-SA 3.0；保留ChEMBL ID、release和来源信息",
    "引用：Mendez et al., Nucleic Acids Research, 2019",
],14.5)
text(s,6.78,1.18,3.4,0.40,"补充来源：项目种子片段",19,True,DARK)
bullets(s,6.80,1.68,5.25,2.15,[
    "10条人工整理的常用小片段",
    "许可：CC0-1.0",
    "用于保留氟、氯、甲基、腈基等基础操作",
    "与ChEMBL工作库统一格式后合并",
],14.5)
# processing flow
flow=[("ChEMBL分子","895,532"),("BRICS单连接点片段","54,570"),("药化过滤后","20,033"),("加入10条seed并去重","20,039")]
for i,(label,count) in enumerate(flow):
    x=0.78+i*3.03
    text(s,x,4.28,2.55,0.34,label,14,True,DARK,align=PP_ALIGN.CENTER)
    text(s,x,4.78,2.55,0.64,count,23,True,BLUE,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE,fill=PALE,border=LINE)
    if i<3:connector(s,x+2.60,5.10,x+2.96,5.10,GRAY,1.2)
text(s,0.80,5.72,11.5,0.76,
     "过滤条件：parent MW ≤ 350；片段重原子 ≤ 12；仅1个连接点；至少2个来源分子；元素限定为C/N/O/S/F/Cl/Br/I；中性、无自由基；去除PAINS和Brenk警示。",
     14,False,BLACK,fill=PALE_BLUE,border=LINE,margin=0.15,valign=MSO_ANCHOR.MIDDLE)

# 12. Docking protocol
s = prs.slides.add_slide(blank); title(s, "08 本次Docking使用的参数", 12)
docking_params = [
    ("程序", "GNINA v1.3.2", "候选和参考均使用同一版本"),
    ("受体", "protein-only receptor PDB", "由完整共晶PDB移除配体后生成"),
    ("配体输入", "candidate SDF / reference-ligand SDF", "候选和参考分别独立对接"),
    ("搜索框", "--autobox_ligand reference-ligand.sdf", "以共晶参考配体定义口袋"),
    ("搜索框扩展", "--autobox_add 4", "在参考配体边界外增加4 Å"),
    ("随机种子", "17、29、43", "每个候选使用三个固定seed"),
    ("输出构象数", "--num_modes 20", "每个seed最多输出20个pose"),
    ("比较pose", "rank-1", "候选与参考按同seed的rank-1 pose比较"),
    ("单次超时", "3600 s", "外部GNINA任务的运行上限"),
]
shape=s.shapes.add_table(len(docking_params)+1,3,Inches(0.72),Inches(1.18),Inches(11.9),Inches(4.42))
tbl=shape.table
for c,v in enumerate(["参数","本次设置","说明"]):tbl.cell(0,c).text=v
tbl.columns[0].width=Inches(2.25);tbl.columns[1].width=Inches(4.15);tbl.columns[2].width=Inches(5.5)
for r,row in enumerate(docking_params,1):
    for c,v in enumerate(row):tbl.cell(r,c).text=v
table_style(tbl,True,10.8)
text(s,0.78,5.88,2.25,0.34,"候选排序规则",16,True,DARK)
ranking = [
    "主指标：minimizedAffinity，越低越好",
    "seed胜率门槛：≥ 2/3",
    "质量惩罚：0.25 × seed标准差",
    "显著改善阈值：0.25；全局尝试上限：80",
]
bullets(s,3.05,5.72,9.25,0.98,ranking,12.8,gap=2)

# 13. Per-site trends
s = prs.slides.add_slide(blank); title(s, "09 代表性运行：分别看每个替换位点的局部趋势", 13)
text(s,0.75,1.05,11.7,0.34,"本次完成78次Docking：Atom 9为30次，Atom 6为38次，Atom 3为10次；2个重复候选未再次对接。",13.5,False,GRAY)
site_specs=[("atom:9","Atom 9",0.48),("atom:6","Atom 6",4.45),("atom:3","Atom 3",8.42)]
for key,name,x in site_specs:
    hs=[v for v in HISTORY if v.get("design_region")==key]
    bestq=None;bestd=None;bests=[];bestitem=None
    for v in hs:
        if v.get("stability_eligible") and v.get("quality") is not None and (bestq is None or v["quality"]>bestq):
            bestq=v["quality"];bestd=v["delta_candidate_minus_reference"];bestitem=v
        bests.append(bestd)
    data=CategoryChartData();data.categories=[str(i+1) for i in range(len(hs))]
    data.add_series("当次候选",[v["delta_candidate_minus_reference"] for v in hs])
    data.add_series("位点best-so-far",bests)
    chart=s.shapes.add_chart(XL_CHART_TYPE.LINE,Inches(x),Inches(1.55),Inches(3.72),Inches(3.85),data).chart
    chart.has_title=True;chart.chart_title.text_frame.text=name
    chart.has_legend=False;chart.value_axis.minimum_scale=-2.1;chart.value_axis.maximum_scale=1.2
    chart.value_axis.has_major_gridlines=True;chart.value_axis.major_gridlines.format.line.color.rgb=LINE
    chart.category_axis.tick_labels.font.size=Pt(8);chart.value_axis.tick_labels.font.size=Pt(8)
    chart.series[0].format.line.color.rgb=RGBColor(150,156,162);chart.series[0].format.line.width=Pt(1.2)
    chart.series[1].format.line.color.rgb=BLUE;chart.series[1].format.line.width=Pt(2.4)
    chart.chart_title.text_frame.paragraphs[0].runs[0].font.name=FONT;chart.chart_title.text_frame.paragraphs[0].runs[0].font.size=Pt(14)
    rmsd=(bestitem.get("pose_consensus") or {}).get("mean_pairwise_rmsd")
    elig=sum(bool(v.get("stability_eligible")) for v in hs)
    note=f"{len(hs)}次；通过seed门槛{elig}次\n最佳Attempt {bestitem['attempt']}，Δ={bestitem['delta_candidate_minus_reference']:.3f}，RMSD={rmsd:.3f} Å"
    text(s,x,5.55,3.72,0.70,note,11.2,True,DARK,align=PP_ALIGN.CENTER,fill=PALE,border=LINE,margin=0.08,valign=MSO_ANCHOR.MIDDLE)
text(s,0.78,6.55,11.65,0.30,"纵轴为 candidate-reference 的 minimizedAffinity；负值表示候选在该协议下更优。",12.5,False,GRAY,align=PP_ALIGN.CENTER)

# 14-19. Every docking attempt, split by site
result_table_specs = [
    ("atom:9", "Atom 9", 15),
    ("atom:6", "Atom 6", 13),
    ("atom:3", "Atom 3", 10),
]
result_page = 14
for site_key, site_name, chunk_size in result_table_specs:
    site_history = [item for item in HISTORY if item.get("design_region") == site_key]
    chunks = [site_history[i:i+chunk_size] for i in range(0, len(site_history), chunk_size)]
    for chunk_index, chunk in enumerate(chunks, 1):
        s = prs.slides.add_slide(blank)
        title(s, f"09 {site_name}：每次Docking的评分结果（{chunk_index}/{len(chunks)}）", result_page)
        text(s,0.55,1.03,12.2,0.34,
             "表中评分均为3个seed的 candidate-reference 平均差：minimizedAffinity负值较好；CNNscore、CNNaffinity和CNN_VS正值较好。",
             11.8,False,GRAY,align=PP_ALIGN.CENTER)
        headers=["Attempt","Parent","Gen","Fragment","Δ minAff","SD","Win","Quality","Δ CNNscore","Δ CNNaff","Δ CNN_VS","Pose RMSD"]
        shape=s.shapes.add_table(len(chunk)+1,len(headers),Inches(0.38),Inches(1.47),Inches(12.55),Inches(5.35))
        tbl=shape.table
        widths=[0.65,0.65,0.50,2.40,1.00,0.80,0.70,1.00,1.30,1.30,1.30,0.95]
        for i,w in enumerate(widths):tbl.columns[i].width=Inches(w)
        for c,v in enumerate(headers):tbl.cell(0,c).text=v
        raw_rows=[]
        for item in chunk:
            metrics=item["comparison"]["metrics"]
            def mean_delta(metric):
                return metrics[metric]["delta_candidate_minus_reference"]["mean"]
            fid=(item.get("transformation") or {}).get("fragment_id") or (item.get("transformation") or {}).get("fragment_smiles") or "-"
            if fid.startswith("chembl-brics-filtered-"):
                fid="ChEMBL-"+fid.rsplit("-",1)[-1]
            parent=item.get("parent_attempt")
            pose=(item.get("pose_consensus") or {}).get("mean_pairwise_rmsd")
            quality=item.get("quality")
            raw_rows.append([
                str(item["attempt"]), str(parent if parent is not None else 0), str(item.get("generation",1)), fid,
                f"{item['delta_candidate_minus_reference']:.3f}", f"{item['seed_stddev']:.3f}",
                f"{item['seed_win_fraction']:.2f}", f"{quality:.3f}" if quality is not None else "—",
                f"{mean_delta('CNNscore'):.3f}", f"{mean_delta('CNNaffinity'):.3f}",
                f"{mean_delta('CNN_VS'):.3f}", f"{pose:.3f}" if pose is not None else "—",
            ])
        for r,row in enumerate(raw_rows,1):
            for c,v in enumerate(row):tbl.cell(r,c).text=v
        table_style(tbl,True,7.8 if len(chunk)>13 else 8.2)
        # Semantic cell coloring: green=favorable, red=unfavorable under the metric direction.
        for r,item in enumerate(chunk,1):
            metrics=item["comparison"]["metrics"]
            values={
                4: item["delta_candidate_minus_reference"] < 0,
                6: bool(item.get("stability_eligible")),
                7: item.get("quality") is not None and item.get("quality") > 0,
                8: metrics["CNNscore"]["delta_candidate_minus_reference"]["mean"] > 0,
                9: metrics["CNNaffinity"]["delta_candidate_minus_reference"]["mean"] > 0,
                10: metrics["CNN_VS"]["delta_candidate_minus_reference"]["mean"] > 0,
                11: ((item.get("pose_consensus") or {}).get("mean_pairwise_rmsd") or 999) <= 2.0,
            }
            for c,is_good in values.items():
                cell=tbl.cell(r,c);cell.fill.solid();cell.fill.fore_color.rgb=PALE_GREEN if is_good else PALE_RED
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:set_run(run,7.8 if len(chunk)>13 else 8.2,False,GREEN if is_good else RED)
        text(s,0.65,6.92,12.0,0.24,"绿色表示该指标方向较好；红色表示该指标方向较差或pose RMSD > 2 Å。",10.5,False,GRAY,align=PP_ALIGN.CENTER)
        result_page += 1

# 20. Best candidate
s = prs.slides.add_slide(blank); title(s, "10 当前最佳候选：Attempt 20", 20)
s.shapes.add_picture(str(ASSETS/"reference_ligand.png"),Inches(0.55),Inches(1.12),width=Inches(3.25),height=Inches(1.68))
s.shapes.add_picture(str(ASSETS/"best_candidate.png"),Inches(3.82),Inches(1.12),width=Inches(3.25),height=Inches(1.68))
text(s,0.72,2.77,6.15,0.28,"Atom 9；generation 3；在parent attempt 15上替换现有取代基",12.5,True,DARK,align=PP_ALIGN.CENTER)
# property table
prop=[("性质","参考","候选"),("MW","323.40","424.50"),("LogP","4.06","3.68"),("HBD/HBA","2/5","2/8"),("TPSA","75.72","97.42"),("重原子","24","31"),("可旋转键","5","7")]
shape=s.shapes.add_table(len(prop),3,Inches(7.35),Inches(1.18),Inches(5.2),Inches(2.02));tbl=shape.table
for r,row in enumerate(prop):
    for c,v in enumerate(row):tbl.cell(r,c).text=v
table_style(tbl,True,10.3)
# seed chart
per=BEST["comparison"]["per_seed"];d=CategoryChartData();d.categories=[str(v["seed"]) for v in per]
d.add_series("候选",[float(v["candidate_pose"]["properties"]["minimizedAffinity"]) for v in per])
d.add_series("参考",[float(v["reference_pose"]["properties"]["minimizedAffinity"]) for v in per])
chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(0.72),Inches(3.30),Inches(6.15),Inches(3.12),d).chart
chart.has_title=True;chart.chart_title.text_frame.text="三个seed的minimizedAffinity"
chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM
chart.series[0].format.fill.solid();chart.series[0].format.fill.fore_color.rgb=BLUE
chart.series[1].format.fill.solid();chart.series[1].format.fill.fore_color.rgb=RGBColor(160,166,172)
chart.value_axis.has_major_gridlines=True;chart.value_axis.major_gridlines.format.line.color.rgb=LINE
chart.chart_title.text_frame.paragraphs[0].runs[0].font.name=FONT;chart.chart_title.text_frame.paragraphs[0].runs[0].font.size=Pt(14)
summary=[("平均相对差","-1.794"),("seed胜率","3/3"),("平均pose RMSD","0.373 Å"),("CNNaffinity均值差","+0.769")]
shape=s.shapes.add_table(5,2,Inches(7.35),Inches(3.48),Inches(5.2),Inches(2.18));tbl=shape.table
tbl.cell(0,0).text="指标";tbl.cell(0,1).text="结果"
for r,row in enumerate(summary,1):
    tbl.cell(r,0).text=row[0];tbl.cell(r,1).text=row[1]
table_style(tbl,True,11.5)
text(s,7.35,5.78,5.18,0.34,"较好的信号：三个seed主指标均改善；pose一致；新增ASN132/GLU12/GLY11接触",10.8,True,GREEN,align=PP_ALIGN.CENTER)
text(s,7.35,6.20,5.18,0.34,"需要注意：seed间幅度有波动；丢失GLU81和LYS89共识接触",10.8,True,RED,align=PP_ALIGN.CENTER)

# 21. Interpretation
s = prs.slides.add_slide(blank); title(s, "11 当前结果的解释和不足", 21)
text(s,0.78,1.25,5.45,0.40,"目前可以说明",19,True,DARK)
bullets(s,0.80,1.78,5.45,3.65,[
    "工作流已完成从结构查询到多seed Docking的闭环运行",
    "Atom 9在本次运行中给出了最好的局部搜索结果",
    "Attempt 20在三个seed上的主指标方向一致，pose也较一致",
    "系统能区分不同位点，并保留各位点和全局历史最佳",
],15.5)
text(s,6.72,1.25,5.45,0.40,"目前仍不能说明",19,True,DARK)
bullets(s,6.74,1.78,5.45,3.65,[
    "运行达到80次硬上限后停止，不代表科学收敛",
    "当前主要是1H1Q单体系，尚无跨靶点验证",
    "刚性受体和Docking打分函数可能带来系统偏差",
    "部分参考接触丢失，且不同CNN指标并非完全一致",
    "没有RBFE和湿实验，不能声称真实活性提升",
],15.5)
text(s,0.82,5.85,11.45,0.62,"当前结论：筛选得到一个值得继续计算验证的候选，而不是已经证明活性提高。",18,True,BLACK,align=PP_ALIGN.CENTER,fill=PALE,border=LINE,margin=0.15,valign=MSO_ANCHOR.MIDDLE)

# 22. Next
s = prs.slides.add_slide(blank); title(s, "12 下一步工作", 22)
nexts=[
    ("1","复核Docking协议","检查参考配体共晶pose重现；对最佳候选做独立重复运行。"),
    ("2","从单一best转向候选集合","按结构簇、位点、分数、seed稳定性和分子性质维护Top-K。"),
    ("3","完成对照实验","比较无工具/有限工具/完整证据门，以及单seed/多seed和parent优化。"),
    ("4","准备RBFE输入","检查共同核心、原子映射、扰动大小、电荷和参数化，不自动启动RBFE。"),
    ("5","扩展测试体系","增加不同靶点和不同配体骨架，评估方法是否具有泛化性。"),
]
for i,(n,t,d) in enumerate(nexts):
    y=1.25+i*1.02
    text(s,0.80,y,0.46,0.46,n,15,True,WHITE,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE,fill=BLUE,border=BLUE)
    text(s,1.50,y-0.01,2.85,0.38,t,17,True,DARK)
    text(s,4.55,y-0.01,7.55,0.46,d,15,False,BLACK)
text(s,0.82,6.55,11.45,0.30,"优先顺序：先验证评分可靠性，再优化搜索策略，最后进入RBFE和实验。",14,True,BLUE,align=PP_ALIGN.CENTER)

# 23. Files
s = prs.slides.add_slide(blank); title(s, "附录：结果和代码位置", 23)
files=[
    ("项目说明","README.md"),
    ("工作流代码","molecular_agent/workflow.py"),
    ("工具实现","molecular_agent/tools.py"),
    ("候选构建","molecular_agent/editing.py"),
    ("本次结果","runs/docking-loop-codex-gpt56-20260902-171109/real/result.json"),
    ("Docking历史",".../real/docking-history.json"),
    ("最佳候选",".../real/candidate-20.sdf"),
    ("参考配体",".../real/reference-ligand.sdf"),
]
shape=s.shapes.add_table(len(files)+1,2,Inches(0.85),Inches(1.30),Inches(11.6),Inches(5.45));tbl=shape.table
tbl.cell(0,0).text="内容";tbl.cell(0,1).text="路径"
tbl.columns[0].width=Inches(2.4);tbl.columns[1].width=Inches(9.2)
for r,row in enumerate(files,1):
    tbl.cell(r,0).text=row[0];tbl.cell(r,1).text=row[1]
table_style(tbl,True,12)

prs.core_properties.title = "蛋白质-配体分子优化工作流：项目进展汇报"
prs.core_properties.subject = "Simple Molecular Agent"
prs.core_properties.comments = "Content-first editable research presentation."
prs.save(OUT)
print(OUT)
