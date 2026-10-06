"""Five-slide, real-4WKQ/Mermaid revision. Previous PPTs are never overwritten."""
from pathlib import Path
import json
import xml.etree.ElementTree as ET
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.opc.package import Part
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
AS=ROOT/'assets/ppt_4wkq'
M=AS/'mermaid'
OUT=ROOT/'deliverables/科研汇报_4WKQ真实结构_Mermaid简洁版.pptx'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
FONT='Microsoft YaHei'
BLUE='204867';INK='2A4051';TEAL='158C99';ORANGE='D89A3B';MUTED='6D8292';LINE='DCE8EE'
WHITE='FFFFFF'

def rgb(c):return RGBColor.from_string(c)

def txt(s,x,y,w,h,value,size=18,color=INK,bold=False,align=PP_ALIGN.LEFT):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,v in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=align
        p.line_spacing=1.10;p.space_after=Pt(4)
        r=p.add_run();r.text=v;r.font.name=FONT;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=rgb(color)
        ea=OxmlElement('a:ea');ea.set('typeface',FONT);r._r.get_or_add_rPr().append(ea)
    return box

def flat(sh):
    for e in sh._element.findall(qn('p:style')):sh._element.remove(e)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh

def line(s,x1,y1,x2,y2,color=LINE,width=1):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=rgb(color);sh.line.width=Pt(width);return flat(sh)

def shape(s,x,y,w,h,color,ellipse=False):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL if ellipse else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background();return flat(sh)

def pic(s,p,x,y,w,h):
    with Image.open(p) as im:iw,ih=im.size
    scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
    return s.shapes.add_picture(str(p),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))

def mermaid(s,name,x,y,w,h):
    svg=M/(name+'.svg');png=M/(name+'.png')
    sh=pic(s,png,x,y,w,h)
    package=s.part.package
    part=Part(package.next_partname('/ppt/media/image%d.svg'),'image/svg+xml',package,svg.read_bytes())
    rid=s.part.relate_to(part,RT.IMAGE)
    blip=sh._pic.blipFill.blip
    extlist=blip.find(qn('a:extLst'))
    if extlist is None:extlist=OxmlElement('a:extLst');blip.append(extlist)
    ext=OxmlElement('a:ext');ext.set('uri','{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    ns='http://schemas.microsoft.com/office/drawing/2016/SVG/main'
    svgb=etree.Element('{'+ns+'}svgBlip',nsmap={'asvg':ns});svgb.set(qn('r:embed'),rid)
    ext.append(svgb);extlist.append(ext)
    sh.name='Mermaid SVG — '+name
    return sh

def footer(s,source='4WKQ · 保留骨架的分子优化智能体'):
    line(s,.60,7.04,12.73,7.04,LINE,.7)
    txt(s,.62,7.17,11.6,.19,source,8.5,MUTED)
    txt(s,12.18,7.11,.53,.25,f'{len(prs.slides):02d}',10,MUTED,align=PP_ALIGN.RIGHT)

def slide(section,title,subtitle=None,source=None):
    s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(WHITE)
    txt(s,.61,.26,11.8,.26,section,10.5,TEAL,True)
    txt(s,.61,.67,12.03,.56,title,28,BLUE,True)
    line(s,.61,1.39,12.73,1.39,LINE,.9)
    if subtitle:txt(s,.64,1.59,12.0,.42,subtitle,16,MUTED)
    footer(s,source or '4WKQ · 保留骨架的分子优化智能体')
    return s

def note(s,t):s.notes_slide.notes_text_frame.text=t

def summary(s,value,y=6.42):
    shape(s,.67,y,.045,.40,TEAL)
    txt(s,.87,y-.01,11.75,.47,value,19.5,BLUE,True)

# 1. Title + actual molecular structure, no invented protein art.
s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(WHITE)
txt(s,.73,.61,6.9,.32,'近期工作进展  /  4WKQ · EGFR–gefitinib',13,TEAL,True)
txt(s,.72,1.65,7.00,1.62,'基于真实结构的\n分子优化智能体',35,BLUE,True)
txt(s,.76,3.57,6.5,.51,'保留骨架 · 侧链探索 · 计算反馈',21,TEAL,True)
txt(s,.77,4.53,6.12,.96,'在已知结合模式上，\n寻找值得进一步验证的分子改造方案。',19,INK)
txt(s,.78,6.33,6.2,.35,'黄婉仪  ·  2026/9/28',13,MUTED)
pic(s,AS/'4wkq_overview.png',7.71,.67,4.95,5.96)
txt(s,8.18,6.57,4.10,.27,'真实晶体结构  |  PDB 4WKQ · 1.85 Å',11,MUTED,align=PP_ALIGN.CENTER)
footer(s,'结构来源：RCSB PDB 4WKQ；链 A，配体 IRE A 1101；PyMOL 按原始坐标渲染')
note(s,'全部三维蛋白、口袋、配体图均由 4WKQ/raw/4WKQ.original.pdb 的原始坐标渲染；未用生成模型重绘，未最小化或修改原子坐标。配体 IRE A 1101，共31个重原子。保留 CSX A 797；显示 altloc blank/A。固定核心碳以青色、可编辑侧链碳以橙色表示，杂原子按元素区分。')

# 2. Overall story: two experimental context panels + one real Mermaid graph.
s=slide('01  /  整体思路','用真实结构约束设计，用计算反馈收敛候选',
        source='结构：PDB 4WKQ（原始坐标）；流程图：Mermaid 源文件 + 内嵌 SVG 矢量图')
txt(s,.86,1.67,4.46,.34,'a  真实复合物',17,BLUE,True)
txt(s,6.48,1.67,5.67,.34,'b  原始结合口袋',17,BLUE,True)
pic(s,AS/'4wkq_overview.png',1.21,2.03,3.35,1.84)
pic(s,AS/'4wkq_pocket_light.png',7.15,2.03,4.04,1.84)
txt(s,1.12,3.92,3.80,.29,'已知配体与结合模式',12.5,MUTED,align=PP_ALIGN.CENTER)
txt(s,7.02,3.92,4.30,.29,'口袋环境为局部改造提供约束',12.5,MUTED,align=PP_ALIGN.CENTER)
mermaid(s,'overview',.71,4.24,11.91,2.14)
summary(s,'核心是“提出 → 验证 → 反馈”，而不是让模型自由生成分子。',6.50)
note(s,'总体框架只保留五个关键节点及反馈回路，不展开各工具内部实现。Mermaid 源文件 assets/ppt_4wkq/mermaid/overview.mmd。结构图 a 为真实 EGFR kinase domain，b 为同一晶体中配体周围6 Å残基表面；使用半透明显示避免遮挡共晶配体，未替换或改变配体姿态。候选排序综合资格与质量函数，不只看单个对接分数。')

# 3. One key loop; only edited chemistry vs real starting structure.
s=slide('02  /  核心策略','骨架不变，侧链受控优化',
        source='原配体与候选结构来自实际 SDF；候选为计算构建结构，不是新的晶体结构')
txt(s,.87,1.65,4.98,.36,'共晶配体  /  gefitinib',18,BLUE,True)
txt(s,7.93,1.65,4.73,.36,'候选示例  /  计算构建',18,BLUE,True)
pic(s,ROOT/'assets/ppt_white/reference_ligand.png',.72,2.06,4.82,1.97)
pic(s,ROOT/'assets/ppt_white/candidate_06.png',7.79,2.06,4.82,1.97)
txt(s,5.66,2.72,1.95,.87,'固定核心\n仅改侧链',18,TEAL,True,PP_ALIGN.CENTER)
shape(s,1.06,4.15,.13,.13,'AEDFE8',True)
txt(s,1.29,4.08,3.30,.3,'浅蓝：保留的核心骨架',12.5,MUTED)
shape(s,7.97,4.15,.13,.13,'FFCB83',True)
txt(s,8.20,4.08,4.15,.3,'浅橙：允许替换的侧链',12.5,MUTED)
mermaid(s,'key_loop',.76,4.48,11.83,1.91)
summary(s,'模型负责提出方案；是否进入候选排序，由计算验证决定。',6.49)
note(s,'参考配体与候选：runs/4wkq-llm-20260927-193317/reference-ligand.sdf 和 candidate-06.sdf。起始配体重原子坐标与 4WKQ 原始晶体坐标匹配。候选 #06 是本次运行最佳，不是共晶结构，也不表示已验证实验活性。高亮按实际21个保留重原子与侧链集合，而非无约束MCS。流程图 key_loop.mmd。细节仅存备注：本次闭集C6切口固定，LLM从78条匿名侧链中选择；初始碰撞延后至docking，最终姿态和碰撞必须通过验证。')

# 4. Only summary evidence, not every candidate / individual tool output.
llm=json.loads((ROOT/'runs/4wkq-llm-20260927-193317/result.json').read_text())
random=json.loads((ROOT/'runs/4wkq-random-baseline-v1/result.json').read_text())
hist=llm['state']['docking_history']
def best_delta(data):
    return next(x['delta_candidate_minus_reference'] for x in data['state']['docking_history'] if x['attempt']==data['result']['best_attempt'])
s=slide('03  /  当前进展','闭环已跑通，策略收益仍需验证','4WKQ 闭集验证：同一候选库，每次运行最多提出 20 个候选。',
        '数据：LLM run 20260927-193317 / random-baseline-v1；沿用原汇报数据快照')
for i,(value,label) in enumerate([
    ('78 / 78','侧链可正确构建'),
    (f"{sum(x['pose_retention']['status']=='passed' for x in hist)} / 20",'LLM 候选通过姿态验证'),
    (f"{sum(bool(x.get('stability_eligible')) for x in hist)} / 20",'LLM 候选具稳定性资格')]):
    x=.89+i*4.15
    txt(s,x,2.28,3.70,.63,value,34,TEAL,True)
    txt(s,x,3.00,3.80,.37,label,17,INK)
    if i<2:line(s,x+3.77,2.35,x+3.77,3.42,LINE,.9)
line(s,.83,3.74,12.48,3.74,LINE,.9)
txt(s,.91,4.04,5.45,.38,'单次最佳 docking 差值',18,BLUE,True)
txt(s,.94,4.61,1.75,.35,'LLM',18,INK)
txt(s,2.73,4.53,2.33,.55,f'{best_delta(llm):.3f}',28,TEAL,True)
txt(s,.94,5.31,1.75,.35,'随机基线',18,INK)
txt(s,2.73,5.23,2.33,.55,f'{best_delta(random):.3f}',28,BLUE,True)
txt(s,.94,6.05,5.47,.3,'Δ = 候选 − 参考；负值更优',13,MUTED)
txt(s,7.24,4.28,5.13,.99,'已证明链路可用，\n尚未证明 LLM 更优。',25,BLUE,True)
txt(s,7.27,5.54,5.14,.63,'单次 docking 结果，\n不能替代重复比较与实验验证。',17,MUTED)
summary(s,'下一步，以等预算重复实验检验真正的筛选收益。',6.50)
note(s,'使用原版两个run的结果快照，不静默替换到其他运行。参考配体3/3 seed校准通过；正文不逐一报告RMSD或20个候选，避免信息过细。78/78是图构建正确，不表示全部初始构象无碰撞；20/20为候选层姿态门控通过；16/20为stability_eligible。主指标minimizedAffinity，展示合格配对seed的平均候选−参考差。LLM best -0.3672533333，随机best -0.5074233333。不得将对接分数当作实验自由能或活性。')

# 5. Short end-to-end takeaway and a three-stage roadmap.
s=slide('04  /  总结与下一步','让 LLM 提出假设，让结构与计算工具把关')
txt(s,.90,2.08,10.95,.87,'目标不是“自动证明一个新药”，\n而是更有依据地缩小候选范围。',26,BLUE,True)
mermaid(s,'next_steps',.88,3.16,11.55,1.82)
line(s,.89,5.16,12.39,5.16,LINE,.9)
for x,head,body in [(.95,'候选结构','可复现的分子改造'),(5.00,'验证证据','结构与计算结果支持'),(9.11,'完整记录','决策和运行过程可追溯')]:
    txt(s,x,5.49,3.27,.39,head,21,TEAL,True)
    txt(s,x,6.03,3.39,.35,body,16,MUTED)
note(s,'下一步：多个LLM/随机/score-guided同预算重复运行；随后按计算证据形成候选清单，并考虑更高精度计算及实验验证。路线图中的实验验证是未来计划，不是已完成结果。')

prs.core_properties.title='基于4WKQ真实结构的分子优化智能体'
prs.core_properties.subject='真实晶体结构 + Mermaid 总体闭环；五页简洁版'
prs.core_properties.author='黄婉仪'
prs.core_properties.comments='3D visuals rendered directly from deposited 4WKQ coordinates. Mermaid SVGs with PNG fallbacks; MMD source supplied. No generative molecular redraw.'
prs.save(OUT)
issues=[]
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if sh.left<0 or sh.top<0 or sh.left+sh.width>prs.slide_width+9144 or sh.top+sh.height>prs.slide_height+9144:
            issues.append((i,sh.name))
assert not issues,issues
(ROOT/'deliverables/4wkq_revision_qc.json').write_text(json.dumps({
    'slide_count':len(prs.slides),'white_backgrounds':5,'mermaid_slides':[2,3,5],
    'embedded_mermaid_svg_count':3,'original_coordinates_unchanged':True,
    'generative_structure_images_used':False,'out_of_bounds_shapes':issues,
    'llm_best_delta':best_delta(llm),'random_best_delta':best_delta(random)
},ensure_ascii=False,indent=2))
print(OUT)
