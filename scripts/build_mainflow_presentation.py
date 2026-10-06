"""Single-topic Mermaid main-flow deck. Never overwrites the previous PPTs."""
from pathlib import Path
import json
import re
import xml.etree.ElementTree as ET

from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.opc.package import Part
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
MMD = ROOT / 'assets/ppt_mainflow/mermaid'
OUT = ROOT / 'deliverables/科研汇报_主流程_Mermaid简洁版.pptx'
FONT = 'Microsoft YaHei'
BLUE, INK, TEAL, ORANGE, MUTED, LINE, WHITE = '204867', '2A4051', '158C99', 'D89A3B', '6D8292', 'DCE8EE', 'FFFFFF'


def rgb(c):
    return RGBColor.from_string(c)


def txt(s, x, y, w, h, value, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, v in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.10
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = v
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        ea = OxmlElement('a:ea')
        ea.set('typeface', FONT)
        r._r.get_or_add_rPr().append(ea)
    return box


def flat(sh):
    for e in sh._element.findall(qn('p:style')):
        sh._element.remove(e)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh


def line(s, x1, y1, x2, y2, color=LINE, width=1):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(width)
    return flat(sh)


def shape(s, x, y, w, h, color, ellipse=False):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL if ellipse else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    return flat(sh)


def pic(s, p, x, y, w, h):
    with Image.open(p) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return s.shapes.add_picture(str(p), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), width=Inches(pw), height=Inches(ph))


def svg_size(p):
    head = p.read_text(encoding='utf-8')[:2000]
    m = re.search(r'viewBox="([^"]+)"', head)
    vb = [float(v) for v in m.group(1).split()]
    return vb[2], vb[3]


def mermaid(s, name, x, y, w):
    svg = MMD / (name + '.svg')
    png = MMD / (name + '.png')
    _, vh = svg_size(svg)
    vw, _ = svg_size(svg)
    h = w * vh / vw
    sh = pic(s, png, x, y, w, h)
    package = s.part.package
    part = Part(package.next_partname('/ppt/media/image%d.svg'), 'image/svg+xml', package, svg.read_bytes())
    rid = s.part.relate_to(part, RT.IMAGE)
    blip = sh._pic.blipFill.blip
    extlist = blip.find(qn('a:extLst'))
    if extlist is None:
        extlist = OxmlElement('a:extLst')
        blip.append(extlist)
    ext = OxmlElement('a:ext')
    ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    ns = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
    svgb = etree.Element('{' + ns + '}svgBlip', nsmap={'asvg': ns})
    svgb.set(qn('r:embed'), rid)
    ext.append(svgb)
    extlist.append(ext)
    sh.name = 'Mermaid SVG — ' + name
    return sh, h


def footer(s, source, page=None):
    line(s, .60, 7.04, 12.73, 7.04, LINE, .7)
    txt(s, .62, 7.17, 11.0, .19, source, 8.5, MUTED)
    if page is not None:
        txt(s, 12.18, 7.11, .53, .25, f'{page:02d}', 10, MUTED, align=PP_ALIGN.RIGHT)


def content_slide(section, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(WHITE)
    txt(s, .61, .26, 11.8, .26, section, 10.5, TEAL, True)
    txt(s, .61, .67, 12.03, .56, title, 28, BLUE, True)
    line(s, .61, 1.39, 12.73, 1.39, LINE, .9)
    if subtitle:
        txt(s, .64, 1.55, 12.0, .40, subtitle, 15.5, MUTED)
    return s


def legend(s, y):
    items = [('AEDFE8', '宿主（确定性计算与校验）'), ('CDEDE8', 'LLM（改造决策）'),
             ('FBE3BE', 'Docking（计算反馈）'), ('E3E8EC', '输入 / 输出')]
    x = .66
    for color, label in items:
        shape(s, x, y + .015, .19, .19, color, ellipse=True)
        box = txt(s, x + .27, y - .015, 2.55, .26, label, 11.5, INK)
        x += .27 + 2.42
    return y


prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)

# ---- Slide 1: cover ----
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb(WHITE)
txt(s, .73, .61, 7.0, .32, '近期工作进展  /  simple_molecular_agent', 13, TEAL, True)
txt(s, .72, 1.63, 8.2, 1.62, '保留骨架的\n分子优化智能体', 35, BLUE, True)
txt(s, .76, 3.55, 7.5, .51, 'LLM 提出假设 · 宿主与计算工具把关', 21, TEAL, True)
txt(s, .77, 4.49, 7.0, .96, '以真实共晶结构为起点，\n通过“提出 → 验证 → 反馈”的闭环收敛候选。', 19, INK)
txt(s, .78, 6.25, 6.2, .35, '黄婉仪  ·  2026/9/28', 13, MUTED)
line(s, 7.95, .95, 7.95, 6.55, LINE, .9)
txt(s, 8.55, 1.72, 4.2, .50, '主流程', 24, BLUE, True)
txt(s, 8.57, 2.22, 4.05, .95, '共晶结构输入\n宿主生成设计 dossier\nLLM 批量规划候选', 15, MUTED)
txt(s, 8.57, 3.28, 4.05, .95, '宿主构建与几何筛选\nscreening docking 打分\n结果反馈并迭代', 15, MUTED)
txt(s, 8.57, 4.34, 4.05, .75, '多 seed 确认\n输出最佳候选 + 完整审计', 15, MUTED)
footer(s, 'simple_molecular_agent · 从完整共晶复合物出发的闭环分子优化', 1)
s.notes_slide.notes_text_frame.text = (
    '本页为封面。完整流程见下一页。项目默认采用批量多位点优化：宿主先准备位点与片段 dossier，'
    'LLM 顺序提出单个单点改造，宿主确定性完成 RDKit 构建、价态/电荷、碰撞、去重与几何检查，'
    '再对几何通过者做 screening docking，并把有界摘要反馈给 LLM；少量 finalist 用多 seed docking 确认，'
    '最终输出历史最佳候选和完整审计。流程图源文件：assets/ppt_mainflow/mermaid/main_flow.mmd。'
)

# ---- Slide 2: main flow ----
s = content_slide('主流程  /  MAIN WORKFLOW', '共晶结构 → 设计 → 构建 → docking → 反馈 → 输出',
                  'LLM 负责提出改造方案，宿主负责确定性校验；docking 结果回传形成闭环。')
legend(s, 2.34)
_, h = mermaid(s, 'main_flow', .50, 2.66, 12.33)
by = 2.66 + h + .30
line(s, .83, by, 12.48, by, LINE, .9)
cards = [
    ('提出', 'LLM 依据配体、口袋与片段证据，批量规划单点改造方案'),
    ('校验', '宿主确定性构建候选并做价态、电荷、碰撞与几何筛选'),
    ('反馈', 'screening docking 打分回传，继续迭代后输出'),
]
for i, (head, body) in enumerate(cards):
    x = .89 + i * 4.15
    txt(s, x, by + .26, 3.35, .38, head, 21, TEAL, True)
    txt(s, x, by + .80, 3.80, .72, body, 14.5, INK)
footer(s, '流程来源：README.md / workflow.mmd；图为 Mermaid 渲染，源文件 assets/ppt_mainflow/mermaid/main_flow.mmd', 2)
s.notes_slide.notes_text_frame.text = (
    '主流程为默认批量优化模式的对外简化表达，只保留六个关键阶段和一条反馈回路，不展开各工具内部实现。\n'
    '输入：完整共晶复合物 PDB + 优化任务；宿主解析蛋白/配体坐标与化学拓扑，一次性生成配体、口袋、相互作用、'
    '合法位点和片段面板的有界 dossier。\n'
    '设计：LLM 按 active target 逐次提出单点改造；宿主精确执行 transformation、RDKit 构建、'
    '价态/电荷、碰撞、去重和几何验证，不把原始坐标或全部片段记录发送给模型。\n'
    '计算：几何通过者用 screening seed 依次 docking，候选与参考配体使用相同 seed 配对打分。\n'
    '反馈与收敛：只把有界摘要（几何通过率、分数、pose、相互作用、best-so-far）回传 LLM；LLM 可继续规划、'
    '刷新片段面板，随后用多 seed docking 确认 seed 稳定性和姿态。\n'
    '停止：LLM 主动停止或达到批次/候选/请求预算；最终输出历史最佳候选和完整审计。'
    'docking 结果是固定协议下的排序信号，不等价于实验活性或真实结合自由能。'
)

assert not OUT.exists(), f'Refusing to overwrite existing deck: {OUT}'
prs.core_properties.title = '保留骨架的分子优化智能体 · 主流程'
prs.core_properties.subject = 'simple_molecular_agent 主流程 Mermaid 流程图'
prs.core_properties.author = '黄婉仪'
prs.core_properties.comments = 'Mermaid main-flow diagram embedded as SVG with PNG fallback; .mmd source supplied.'
prs.save(OUT)

issues = []
for i, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if sh.left < 0 or sh.top < 0 or sh.left + sh.width > prs.slide_width + 9144 or sh.top + sh.height > prs.slide_height + 9144:
            issues.append((i, sh.name))
assert not issues, issues
(ROOT / 'deliverables/mainflow_revision_qc.json').write_text(json.dumps({
    'slide_count': len(prs.slides),
    'mermaid_slides': [2],
    'embedded_mermaid_svg_count': 1,
    'mermaid_source': 'assets/ppt_mainflow/mermaid/main_flow.mmd',
    'out_of_bounds_shapes': issues,
}, ensure_ascii=False, indent=2))
print(OUT)
