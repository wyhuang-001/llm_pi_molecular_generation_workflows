"""Main-flow deck (revision 2).

Content changes requested:
- the input block now describes the LLM input context (task, ligand graph,
  pocket/interaction, legal sites, fragment library) and shows the system-side
  structure parsing explicitly;
- the LLM step reads "提出改造思路" (batch planning removed);
- the word "宿主" is not used anywhere;
- docking is followed by an RBFE stage;
- Playwright MCP external research is part of the input side;
- the feedback arrow states that results are written back to update the context.

Previous PPTs are never overwritten.
"""
from pathlib import Path
import json
import re

from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.opc.package import Part
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
MMD = ROOT / 'assets/ppt_mainflow/mermaid'
OUT = ROOT / 'deliverables/科研汇报_主流程_Mermaid简洁版_修订版.pptx'
FONT = 'Microsoft YaHei'
BLUE, INK, TEAL, ORANGE, MUTED, LINE, WHITE = '204867', '2A4051', '158C99', 'D89A3B', '6D8292', 'DCE8EE', 'FFFFFF'


def rgb(c):
    return RGBColor.from_string(c)


def txt(s, x, y, w, h, value, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, v in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.10
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = v
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        ea = OxmlElement('a:ea')
        ea.set('typeface', FONT)
        r._r.get_or_add_rPr().append(ea)
    return box


def flat(sh):
    for e in sh._element.findall(qn('p:style')):
        sh._element.remove(e)
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    return sh


def line(s, x1, y1, x2, y2, color=LINE, width=1):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(width)
    return flat(sh)


def shape(s, x, y, w, h, color, ellipse=False):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL if ellipse else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    return flat(sh)


def pic(s, p, x, y, w, h):
    with Image.open(p) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return s.shapes.add_picture(str(p), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), width=Inches(pw), height=Inches(ph))


def svg_size(p):
    head = p.read_text(encoding='utf-8')[:2000]
    vb = [float(v) for v in re.search(r'viewBox="([^"]+)"', head).group(1).split()]
    return vb[2], vb[3]


def mermaid(s, name, x, y, w):
    svg = MMD / (name + '.svg')
    png = MMD / (name + '.png')
    vw, vh = svg_size(svg)
    h = w * vh / vw
    sh = pic(s, png, x, y, w, h)
    package = s.part.package
    part = Part(package.next_partname('/ppt/media/image%d.svg'), 'image/svg+xml', package, svg.read_bytes())
    rid = s.part.relate_to(part, RT.IMAGE)
    blip = sh._pic.blipFill.blip
    extlist = blip.find(qn('a:extLst'))
    if extlist is None:
        extlist = OxmlElement('a:extLst')
        blip.append(extlist)
    ext = OxmlElement('a:ext')
    ext.set('uri', '{96DAC541-7B7A-43D3-8B79-37D633B846F1}')
    ns = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'
    svgb = etree.Element('{' + ns + '}svgBlip', nsmap={'asvg': ns})
    svgb.set(qn('r:embed'), rid)
    ext.append(svgb)
    extlist.append(ext)
    sh.name = 'Mermaid SVG — ' + name
    return sh, h


def footer(s, source, page=None):
    line(s, .60, 7.04, 12.73, 7.04, LINE, .7)
    txt(s, .62, 7.17, 11.0, .19, source, 8.5, MUTED)
    if page is not None:
        txt(s, 12.18, 7.11, .53, .25, f'{page:02d}', 10, MUTED, align=PP_ALIGN.RIGHT)


def content_slide(section, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(WHITE)
    txt(s, .61, .26, 11.8, .26, section, 10.5, TEAL, True)
    txt(s, .61, .67, 12.03, .56, title, 27, BLUE, True)
    line(s, .61, 1.39, 12.73, 1.39, LINE, .9)
    if subtitle:
        txt(s, .64, 1.52, 12.0, .40, subtitle, 15, MUTED)
    return s


prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)

# ---- Slide 1: cover ----
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = rgb(WHITE)
txt(s, .73, .61, 7.0, .32, '近期工作进展  /  simple_molecular_agent', 13, TEAL, True)
txt(s, .72, 1.63, 8.2, 1.62, '保留骨架的\n分子优化智能体', 35, BLUE, True)
txt(s, .76, 3.55, 7.5, .51, 'LLM 提出改造思路 · 系统与计算工具把关', 20, TEAL, True)
txt(s, .77, 4.49, 7.0, .96, '以真实共晶结构为起点，系统组装上下文，\nLLM 提出思路，docking 与 RBFE 结果回写并迭代。', 18, INK)
txt(s, .78, 6.25, 6.2, .35, '黄婉仪  ·  2026/9/28', 13, MUTED)
line(s, 7.95, .95, 7.95, 6.55, LINE, .9)
txt(s, 8.55, 1.72, 4.2, .50, '主流程', 24, BLUE, True)
txt(s, 8.57, 2.28, 4.05, .95, '任务 + 共晶复合物 PDB\n系统解析结构 · 组装上下文\nPlaywright MCP 外部检索', 14, MUTED)
txt(s, 8.57, 3.42, 4.05, .95, 'LLM 提出改造思路\n系统构建候选并校验\ndocking + RBFE 计算', 14, MUTED)
txt(s, 8.57, 4.56, 4.05, .75, '结果回写，更新输入上下文\n输出最佳候选 + 完整审计', 14, MUTED)
footer(s, 'simple_molecular_agent · 从完整共晶复合物出发的闭环分子优化', 1)
s.notes_slide.notes_text_frame.text = (
    '本页为封面，完整流程见下一页。系统先解析共晶结构，恢复配体分子图、口袋残基与相互作用、合法位点，'
    '并把任务、片段库以及 Playwright MCP 的一次性外部检索结果一起组织为 LLM 的输入上下文；'
    'LLM 在此上下文下提出改造思路，系统完成候选构建与化学/几何校验，随后进行 docking 与 RBFE 计算，'
    '结果回写并更新上下文，驱动下一轮搜索。'
)

# ---- Slide 2: main flow ----
s = content_slide('主流程  /  MAIN WORKFLOW', '输入上下文 → 提出思路 → 计算 → 反馈',
                  '系统解析结构并组织上下文；LLM 提出改造思路；docking 与 RBFE 结果回写上下文。')
mermaid(s, 'main_flow', .50, 1.72, 8.10)

# right column: legend + key points
txt(s, 8.92, 1.74, 3.80, .30, '图例', 14, BLUE, True)
legend_items = [
    ('D9E1F5', '输入与上下文'),
    ('CFE3EE', '系统：解析结构、构建与校验'),
    ('B9E2DE', 'LLM：提出改造思路'),
    ('F7D9A8', '计算反馈：docking / RBFE'),
]
for i, (color, label) in enumerate(legend_items):
    y = 2.16 + i * .34
    shape(s, 8.92, y + .02, .17, .17, color, ellipse=True)
    txt(s, 9.18, y - .02, 3.45, .30, label, 11.5, INK)
line(s, 8.92, 3.68, 12.72, 3.68, LINE, .9)
txt(s, 8.92, 3.84, 3.80, .30, '要点', 14, BLUE, True)
points = [
    '上下文包含任务、配体分子图、口袋与相互作用、合法位点与片段库。',
    'LLM 只提出改造思路；候选构建与化学、几何校验由系统完成。',
    'docking 与 RBFE 结果回写并更新上下文，驱动下一轮搜索。',
]
for i, body in enumerate(points):
    y = 4.24 + i * .66
    shape(s, 8.94, y + .05, .045, .30, TEAL)
    txt(s, 9.12, y, 3.52, .62, body, 11.5, INK)
txt(s, 8.94, 6.44, 3.74, .48, '计算结果是固定协议下的筛选信号，\n不等同于实验活性或真实结合自由能。', 10.5, MUTED)
footer(s, '流程来源：README.md / workflow.mmd；图为 Mermaid 渲染，源文件 assets/ppt_mainflow/mermaid/main_flow.mmd', 2)
s.notes_slide.notes_text_frame.text = (
    '本页是整体主流程，只保留关键阶段与一条反馈回路，不展开各工具内部实现。\n'
    '输入与上下文：任务与优化目标、完整共晶复合物 PDB、以及可选的 Playwright MCP 一次性外部检索（结构页面、截图与来源）。'
    '系统解析 PDB，恢复配体分子图、口袋残基与相互作用、合法位点，并结合片段库组装为 LLM 的输入上下文；'
    '原始坐标与全部片段记录不直接发送给模型，只发送有界摘要。\n'
    '提出思路：LLM 在此上下文下提出改造思路（改哪个位点、接什么片段、依据与风险）。\n'
    '构建与校验：系统精确执行分子图编辑，完成价态/电荷、碰撞、去重与几何校验。\n'
    '计算：几何通过者先做 screening docking（多 seed，与参考配体同 seed 配对，并分析相互作用），'
    '少量 finalist 再做更高精度的 RBFE 计算。\n'
    '反馈与输出：docking 与 RBFE 的结果（分数、pose、相互作用、失败原因）回写并更新输入上下文，'
    '驱动 LLM 的下一轮思路；满足停止条件后输出最佳候选与完整审计。'
)

assert not OUT.exists(), f'Refusing to overwrite existing deck: {OUT}'
prs.core_properties.title = '保留骨架的分子优化智能体 · 主流程（修订版）'
prs.core_properties.subject = 'simple_molecular_agent 主流程 Mermaid 流程图：输入上下文、LLM 思路、docking 与 RBFE、反馈更新上下文'
prs.core_properties.author = '黄婉仪'
prs.core_properties.comments = 'Mermaid main-flow diagram embedded as SVG with PNG fallback; .mmd source supplied.'
prs.save(OUT)

issues = []
for i, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        if sh.left < 0 or sh.top < 0 or sh.left + sh.width > prs.slide_width + 9144 or sh.top + sh.height > prs.slide_height + 9144:
            issues.append((i, sh.name))
assert not issues, issues
(ROOT / 'deliverables/mainflow_revision_v2_qc.json').write_text(json.dumps({
    'slide_count': len(prs.slides),
    'mermaid_slides': [2],
    'embedded_mermaid_svg_count': 1,
    'mermaid_source': 'assets/ppt_mainflow/mermaid/main_flow.mmd',
    'out_of_bounds_shapes': issues,
}, ensure_ascii=False, indent=2))
print(OUT)
